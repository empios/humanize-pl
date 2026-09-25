"""Structure-aware DOCX traversal and conservative in-place text editing."""

from __future__ import annotations

import hashlib
import json
import shutil
from collections.abc import Iterator
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any
from zipfile import ZipFile


@dataclass
class TextUnit:
    paragraph: Any
    location: str
    text: str
    protected: bool = False
    protection_reasons: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class DocxInventory:
    parts: tuple[str, ...]
    tables: int
    paragraphs: int
    sections: int
    drawings: int
    hyperlinks: int
    bookmarks: int
    fields: int
    comments: int
    tracked_changes: int
    content_controls: int
    text_boxes: int
    images_without_alt: int
    header_footer_text: tuple[str, ...]
    notes_text: tuple[str, ...]
    integrity: dict[str, str] = field(default_factory=dict)
    paragraph_formats: tuple[str, ...] = ()
    document_format: str = ""

    def structural_differences(
        self, other: DocxInventory, *, expected_paragraph_delta: int = 0,
        allow_formatting_changes: bool = False,
    ) -> list[str]:
        """Structural drift between two inventories, ignoring a declared change.

        `expected_paragraph_delta` is how many paragraphs the caller said it
        was going to add, and it is checked for equality rather than as a
        ceiling. "Paragraphs may grow" would let a rewrite drop one clause
        while adding another and still pass; "paragraphs must grow by exactly
        three" fails the moment anything else moves. Every other count stays
        exact, so this widens the guard in one dimension by one declared
        amount and nowhere else.
        """
        differences: list[str] = []
        if expected_paragraph_delta:
            observed = other.paragraphs - self.paragraphs
            if observed != expected_paragraph_delta:
                differences.append(
                    f"paragraphs: {self.paragraphs} → {other.paragraphs} "
                    f"(zapowiedziano {expected_paragraph_delta:+d}, "
                    f"jest {observed:+d})"
                )
        comparable = (
            "tables",
            *(() if expected_paragraph_delta else ("paragraphs",)),
            "sections",
            "drawings",
            "hyperlinks",
            "bookmarks",
            "fields",
            "comments",
            "tracked_changes",
            "content_controls",
            "text_boxes",
            "images_without_alt",
        )
        for name in comparable:
            if getattr(self, name) != getattr(other, name):
                differences.append(
                    f"{name}: {getattr(self, name)} → {getattr(other, name)}"
                )
        if self.header_footer_text != other.header_footer_text:
            differences.append("zmieniła się treść nagłówków lub stopek")
        if self.notes_text != other.notes_text:
            differences.append("zmieniła się treść przypisów lub komentarzy")
        # Style/theme parts may be introduced by normalization or a template,
        # so compare only semantic package parts that must never disappear.
        protected_parts = {
            name
            for name in self.parts
            if name.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes"))
            or name in {"word/comments.xml"}
        }
        if not protected_parts.issubset(set(other.parts)):
            differences.append("z pakietu zniknęła chroniona część OOXML")
        for label, digest in self.integrity.items():
            if other.integrity.get(label) != digest:
                differences.append(f"zmieniono chronioną strukturę lub treść: {label}")
        if not allow_formatting_changes:
            # Declared drafts may add paragraph styles, never remove or reorder old ones.
            remaining = iter(other.paragraph_formats)
            if any(not any(candidate == original for candidate in remaining) for original in self.paragraph_formats):
                differences.append("zmieniono formatowanie istniejących akapitów lub runów")
            if not expected_paragraph_delta and self.paragraph_formats != other.paragraph_formats:
                differences.append("zmieniono układ formatowania akapitów")
            if self.document_format != other.document_format:
                differences.append("zmieniono style lub formatowanie dokumentu w trybie zachowania")
        return differences


def _paragraph_text(paragraph: Any) -> str:
    # One analysis line per DOCX paragraph, including inline separators.
    # U+2028 preserves a manual break without shifting paragraph indices.
    from docx.oxml.ns import qn

    separators = {qn("w:tab"): "\t", qn("w:br"): "\u2028", qn("w:cr"): "\u2028"}
    return "".join(
        separators.get(node.tag, (node.text or "").replace("\n", "\u2028"))
        for node in paragraph._p.xpath(".//w:t | .//w:tab | .//w:br | .//w:cr")
    )


def _protection_reasons(paragraph: Any) -> list[str]:
    checks = {
        "pole dokumentu": ".//w:fldChar | .//w:instrText | .//w:fldSimple",
        "komentarz": ".//w:commentRangeStart | .//w:commentRangeEnd | .//w:commentReference",
        "śledzona zmiana": ".//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo",
        "kontrolka treści": ".//w:sdt",
        "pole tekstowe": ".//w:txbxContent",
        "hiperłącze": ".//w:hyperlink",
        "zakładka / cel odesłania": ".//w:bookmarkStart | .//w:bookmarkEnd",
        "odnośnik przypisu": ".//w:footnoteReference | .//w:endnoteReference",
        "podział lub tabulator": ".//w:br | .//w:cr | .//w:tab",
        "obiekt graficzny": ".//w:drawing | .//w:pict | .//w:object",
    }
    return [label for label, xpath in checks.items() if paragraph._p.xpath(xpath)]


def iter_text_units(document: Any) -> Iterator[TextUnit]:
    """Walk actual XML cells and wrappers in order, preserving spanning ranges."""
    from docx.oxml.ns import qn
    from docx.oxml.table import CT_Tbl  # type: ignore
    from docx.oxml.text.paragraph import CT_P  # type: ignore
    from docx.table import Table, _Cell  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore

    ranges: dict[str, set[str]] = {"comment": set(), "bookmark": set()}
    field_depth = 0

    def walk(container, parent, location, inherited=()):
        nonlocal field_depth
        paragraph_index = table_index = 0
        for child_index, child in enumerate(container.iterchildren(), 1):
            if isinstance(child, CT_P):
                paragraph_index += 1
                paragraph = Paragraph(child, parent)
                reasons = list(inherited) + _protection_reasons(paragraph)
                if any(ranges.values()):
                    reasons.append("zakres komentarza lub zakładki")
                if field_depth:
                    reasons.append("wieloakapitowe pole dokumentu")
                for node in child.iter():
                    for key, start, end in (
                        ("comment", "commentRangeStart", "commentRangeEnd"),
                        ("bookmark", "bookmarkStart", "bookmarkEnd"),
                    ):
                        if node.tag == qn(f"w:{start}"):
                            ranges[key].add(node.get(qn("w:id"), ""))
                        elif node.tag == qn(f"w:{end}"):
                            ranges[key].discard(node.get(qn("w:id"), ""))
                    if node.tag == qn("w:fldChar"):
                        kind = node.get(qn("w:fldCharType"))
                        field_depth = field_depth + 1 if kind == "begin" else max(0, field_depth - (kind == "end"))
                text = _paragraph_text(paragraph)
                if text.strip():
                    yield TextUnit(paragraph, f"{location}/paragraph:{paragraph_index}", text,
                                   bool(reasons), list(dict.fromkeys(reasons)))
            elif isinstance(child, CT_Tbl):
                table_index += 1
                table = Table(child, parent)
                for row_index, row in enumerate(child.tr_lst, 1):
                    for cell_index, cell in enumerate(row.tc_lst, 1):
                        yield from walk(cell, _Cell(cell, table),
                                        f"{location}/table:{table_index}/row:{row_index}/cell:{cell_index}", inherited)
            elif child.find(".//" + qn("w:p")) is not None:
                # Block controls / revisions are visible to analysis but never edited.
                label = "kontrolka treści" if child.tag == qn("w:sdt") else "chroniony kontener OOXML"
                yield from walk(child, parent, f"{location}/{child.tag.split('}')[-1]}:{child_index}", (*inherited, label))

    yield from walk(document.element.body, document, "body")


def replace_unit_text(unit: TextUnit, replacement: str) -> None:
    """Replace only changed text nodes; keep runs, links and XML anchors."""
    if unit.protected:
        raise ValueError(f"Chroniony fragment DOCX nie może być redagowany: {unit.location}")
    nodes = list(unit.paragraph._p.xpath(".//w:t"))
    original = "".join(node.text or "" for node in nodes)
    if original == replacement:
        return
    if not nodes:
        raise ValueError(f"Brak węzłów tekstowych w {unit.location}")

    positions: list[tuple[int, int]] = []
    cursor = 0
    for node in nodes:
        value = node.text or ""
        positions.append((cursor, cursor + len(value)))
        cursor += len(value)

    pieces: list[list[str]] = [[] for _ in nodes]
    for operation, start, end, new_start, new_end in SequenceMatcher(
        a=original, b=replacement, autojunk=False,
    ).get_opcodes():
        if operation == "equal":
            for index, (left, right) in enumerate(positions):
                if max(left, start) < min(right, end):
                    pieces[index].append(original[max(left, start):min(right, end)])
        elif operation in {"replace", "insert"}:
            owner = start
            # Removing an introductory run and capitalising the remaining word
            # must take that word's formatting, not the deleted introduction's.
            added = replacement[new_start:new_end]
            if operation == "replace" and len(added) == 1 and original[end - 1:end].casefold() == added.casefold():
                owner = end - 1
            index = next((i for i, (left, right) in enumerate(positions) if left <= owner < right), len(nodes) - 1)
            pieces[index].append(replacement[new_start:new_end])
    from docx.oxml.ns import qn

    for node, parts in zip(nodes, pieces):
        node.text = "".join(parts)
        if node.text and (node.text[0].isspace() or node.text[-1].isspace()):
            node.set(qn("xml:space"), "preserve")

    unit.text = replacement


# Copied properties that would add something the inventory counts or change
# what is around the new paragraph: a section break carried in `w:pPr`, a
# tracked-change mark, and list numbering - a paragraph that joined an
# automatic "§ 1, § 2, ..." list would renumber every unit below it, the one
# thing a "§ 2a" heading exists to avoid.
_NOT_COPIED = (
    "w:sectPr",
    "w:numPr",
    "w:pPrChange",
    "w:rPrChange",
    "w:ins",
    "w:del",
    "w:moveFrom",
    "w:moveTo",
)


def new_paragraph_like(model: Any, text: str) -> Any:
    """A new `w:p` carrying `text` in the formatting of the `model` paragraph.

    Only paragraph properties and the first run's character properties are
    copied - enough for the clause to look like its neighbours, and nothing
    that belongs to the neighbour itself (bookmarks, fields, comments).
    """
    from copy import deepcopy

    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore

    def cleaned(element: Any) -> Any:
        copy = deepcopy(element)
        for name in _NOT_COPIED:
            for node in copy.findall(".//" + qn(name)):
                node.getparent().remove(node)
        return copy

    paragraph = OxmlElement("w:p")
    properties = model._p.find(qn("w:pPr"))
    if properties is not None:
        paragraph.append(cleaned(properties))
    # A copied paragraph style can inherit list numbering even after removing numPr.
    # Explicit numId=0 keeps a new clause from renumbering existing list items.
    ppr = paragraph.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    numbering.append(num_id)
    ppr.append(numbering)
    run = OxmlElement("w:r")
    first_run = model._p.find(".//" + qn("w:r"))
    run_properties = first_run.find(qn("w:rPr")) if first_run is not None else None
    if run_properties is not None:
        run.append(cleaned(run_properties))
    node = OxmlElement("w:t")
    node.text = text
    node.set(qn("xml:space"), "preserve")
    run.append(node)
    paragraph.append(run)
    return paragraph


def document_text(document: Any, *, include_protected: bool = True) -> str:
    return "\n".join(
        unit.text for unit in iter_text_units(document) if include_protected or not unit.protected
    )


def load_document(path: str | Path) -> Any:
    from docx import Document  # type: ignore

    return Document(str(Path(path)))


def inventory_docx(path: str | Path) -> DocxInventory:
    path = Path(path)
    with ZipFile(path) as archive:
        names = tuple(sorted(archive.namelist()))
        xml_parts = {
            name: archive.read(name)
            for name in names
            if name.endswith(".xml") and name.startswith("word/")
        }
        relationships = {name: archive.read(name) for name in names if name.endswith(".rels")}
        binaries = {name: hashlib.sha256(archive.read(name)).hexdigest()
                    for name in names if name.startswith(("word/media/", "word/embeddings/"))}
    from lxml import etree

    namespace = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    }
    roots = {}
    for name, value in xml_parts.items():
        try:
            roots[name] = etree.fromstring(value)
        except etree.XMLSyntaxError:
            continue

    def count(xpath: str) -> int:
        return sum(len(root.xpath(xpath, namespaces=namespace)) for root in roots.values())

    def texts(prefixes: tuple[str, ...]) -> tuple[str, ...]:
        rows: list[str] = []
        for name in sorted(roots):
            if name.startswith(prefixes):
                rows.append("".join(roots[name].xpath(".//w:t/text()", namespaces=namespace)))
        return tuple(rows)

    image_properties = [
        prop
        for root in roots.values()
        for prop in root.xpath(".//wp:docPr", namespaces=namespace)
    ]
    images_without_alt = sum(
        1 for prop in image_properties if not (prop.get("descr") or prop.get("title"))
    )

    main = roots.get("word/document.xml")

    def signature(node):
        # Expanded tag names avoid namespace-prefix changes during serialization.
        return (node.tag, sorted(node.attrib.items()),
                node.text if node.tag.endswith(("}t", "}delText", "}instrText")) else (node.text or "").strip(),
                [signature(child) for child in node if isinstance(child.tag, str)])

    def digest(value):
        return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()

    def selected(xpath):
        return [(name, [signature(node) for node in root.xpath(xpath, namespaces=namespace)])
                for name, root in sorted(roots.items()) if name not in {"word/styles.xml", "word/stylesWithEffects.xml"}]

    integrity = {
        "relacje i cele hiperłączy": digest([(name, signature(etree.fromstring(data))) for name, data in sorted(relationships.items())]),
        "media i osadzone pliki": digest(binaries),
        "nagłówki, stopki, przypisy i komentarze": digest([
            (name, signature(root)) for name, root in sorted(roots.items())
            if name.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes", "word/comments"))
        ]),
        "definicje numeracji": digest(signature(roots["word/numbering.xml"]) if "word/numbering.xml" in roots else None),
        "numeracja akapitów": digest(selected(".//w:numPr[not(w:numId/@w:val='0')]")),
        "powiązania numeracji ze stylami": digest([
            (node.get(f"{{{namespace['w']}}}styleId"),
             [signature(part) for part in node.xpath("./w:basedOn | .//w:numPr", namespaces=namespace)])
            for node in roots.get("word/styles.xml", []) if node.xpath("./w:basedOn | .//w:numPr", namespaces=namespace)
        ]),
        "kotwice, pola, hiperłącza i elementy chronione": digest(selected(
            ".//w:hyperlink | .//w:bookmarkStart | .//w:bookmarkEnd | .//w:fldChar | .//w:instrText | .//w:fldSimple"
            " | .//w:commentRangeStart | .//w:commentRangeEnd | .//w:commentReference"
            " | .//w:footnoteReference | .//w:endnoteReference | .//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo"
            " | .//w:sdt | .//w:drawing | .//w:pict | .//w:object | .//w:br | .//w:cr | .//w:tab"
        )),
        "siatka i scalenia tabel": digest(selected(".//w:tblGrid | .//w:gridSpan | .//w:vMerge | .//w:hMerge")),
    }
    paragraph_formats = tuple(digest([
        signature(node) for node in paragraph.xpath("./w:pPr | ./w:r/w:rPr", namespaces=namespace)
    ]) for paragraph in main.xpath(".//w:p", namespaces=namespace)) if main is not None else ()
    document_format = digest([
        (name, signature(roots[name])) for name in ("word/styles.xml", "word/theme/theme1.xml", "word/fontTable.xml") if name in roots
    ] + selected(".//w:sectPr | .//w:tblPr | .//w:trPr | .//w:tcPr"))

    def main_count(xpath: str) -> int:
        return len(main.xpath(xpath, namespaces=namespace)) if main is not None else 0
    return DocxInventory(
        parts=names,
        tables=main_count(".//w:tbl"),
        paragraphs=main_count(".//w:p"),
        sections=main_count(".//w:sectPr"),
        drawings=count(".//w:drawing | .//w:pict"),
        hyperlinks=count(".//w:hyperlink"),
        bookmarks=count(".//w:bookmarkStart"),
        fields=count(".//w:fldChar | .//w:instrText | .//w:fldSimple"),
        comments=count(".//w:commentRangeStart | .//w:commentReference"),
        tracked_changes=count(".//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo"),
        content_controls=count(".//w:sdt"),
        text_boxes=count(".//w:txbxContent"),
        images_without_alt=images_without_alt,
        header_footer_text=texts(("word/header", "word/footer")),
        notes_text=texts(("word/footnotes", "word/endnotes", "word/comments")),
        integrity=integrity,
        paragraph_formats=paragraph_formats,
        document_format=document_format,
    )


def save_with_inventory_guard(
    document: Any,
    source: Path,
    target: Path,
    *,
    expected_paragraph_delta: int = 0,
    allow_formatting_changes: bool = False,
) -> list[str]:
    """Save a candidate, reverting to a source copy on structural mismatch.

    `expected_paragraph_delta` lets a caller that deliberately adds
    paragraphs - a missing clause the document owed its category - say so in
    advance. It is an exact figure, not a licence to grow: everything else is
    still compared for equality, and a run that adds three paragraphs while
    losing one fails as loudly as before.
    """
    from humanize_pl.io.atomic import atomic_output

    before = inventory_docx(source)
    with atomic_output(target, sources=[source]) as staged:
        document.save(str(staged))
        after = inventory_docx(staged)
        differences = before.structural_differences(
            after, expected_paragraph_delta=expected_paragraph_delta,
            allow_formatting_changes=allow_formatting_changes,
        )
        if differences:
            shutil.copyfile(source, staged)
    return differences
