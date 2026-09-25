"""Review final proposals against immutable source paragraphs.

Edits affecting the same paragraph are reviewed together: accepting a deletion
without its dependent rewrite could change meaning. Rebuilding always starts
from the source, never from a previously reviewed result.
"""

from __future__ import annotations

import hashlib
import json
from copy import copy, deepcopy
from pathlib import Path

from humanize_pl.document import DocumentType, FormatPolicy
from humanize_pl.drafting import drafts_from_payload, insert_drafts, inserted_line_indices
from humanize_pl.io.atomic import atomic_output, ensure_distinct_paths, write_text_atomic


def digest(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _nli_context(source: str, output: str, *, requested: bool, before: dict, after: dict) -> dict:
    assessments = {}
    for text, report in ((source, before), (output, after)):
        if report:
            assessments[digest(text)] = deepcopy(report)
    return {"requested": requested or bool(assessments), "assessments": assessments}


def create_review(source: str, outcome, *, settings=None) -> dict:
    drafts = [dict(row) for row in outcome.drafted_sections if row.get("inserted", True)]
    output = outcome.text_out if outcome.text_out is not None else source
    added = set(inserted_line_indices(drafts_from_payload(drafts)))
    edited = [line for i, line in enumerate(output.split("\n")) if i not in added]
    original = source.split("\n")
    if len(original) != len(edited):
        return {"available": False, "reason": "Zmiana podziału akapitów wymaga przeglądu całego dokumentu."}
    proposals = []
    for i, (before, after) in enumerate(zip(original, edited)):
        if before == after:
            continue
        proposals.append({"id": f"p{i + 1}-{digest(before + chr(0) + after)[:12]}",
                          "kind": "edit", "paragraph_index": i, "before": before, "after": after,
                          "previous": original[i - 1] if i else "",
                          "following": original[i + 1] if i + 1 < len(original) else ""})
    for i, draft in enumerate(drafts):
        proposals.append({"id": f"d{i + 1}-{digest(json.dumps(draft, sort_keys=True))[:12]}",
                          "kind": "draft", "before": "", "after": "\n".join(drafts_from_payload([draft])[0].lines),
                          "draft": draft})
    return {"schema": 1, "available": True, "source": source, "source_sha256": digest(source),
            "document_type": outcome.document_type, "general_options": outcome.general_options,
            "proposals": proposals,
            "nli": _nli_context(
                source, output,
                requested=outcome.requested_operations.get("nli", False),
                before=outcome.nli_before, after=outcome.nli_after,
            ),
            "analysis_settings": ({"check_completeness": settings.check_completeness,
                                   "blueprint": str(settings.blueprint) if settings.blueprint is not None else None,
                                   "style_profile": str(settings.style_profile) if settings.style_profile is not None else None}
                                  if settings is not None else {})}


def select_review(plan: dict, accepted: list[str]) -> tuple[str, list[dict]]:
    if not plan.get("available") or plan.get("schema") != 1:
        raise ValueError("Ten raport nie zawiera obsługiwanego przeglądu propozycji.")
    source = plan["source"]
    if digest(source) != plan.get("source_sha256"):
        raise ValueError("Źródło przeglądu nie zgadza się z zapisanym odciskiem.")
    rows = plan["proposals"]
    ids = [row["id"] for row in rows]
    if len(set(ids)) != len(ids) or not set(accepted) <= set(ids) or len(set(accepted)) != len(accepted):
        raise ValueError("Nieznane lub powtórzone identyfikatory propozycji.")
    lines = source.split("\n")
    touched, drafts = set(), []
    for row in rows:
        if row["kind"] == "edit":
            i = row["paragraph_index"]
            if not isinstance(i, int) or not 0 <= i < len(lines) or i in touched or lines[i] != row["before"]:
                raise ValueError("Propozycja nie pasuje do akapitu źródłowego.")
            if "\n" in row["after"]:
                raise ValueError("Propozycja zmienia podział akapitów.")
            touched.add(i)
            if row["id"] in accepted:
                lines[i] = row["after"]
        elif row["kind"] == "draft":
            if row["id"] in accepted:
                drafts.append(row["draft"])
        else:
            raise ValueError("Nieznany rodzaj propozycji.")
    return insert_drafts("\n".join(lines), drafts_from_payload(drafts)), drafts


def _restore_nli(plan: dict, text: str, item) -> None:
    """Reuse only an assessment of these exact bytes; never infer coverage."""
    context = plan.get("nli") or {}
    if not context.get("requested"):
        return
    item.requested_operations["nli"] = True
    assessments = context.get("assessments", {})
    item.nli_before = deepcopy(assessments.get(plan["source_sha256"], {}))
    assessment = assessments.get(digest(text))
    if assessment:
        item.nli_after = deepcopy(assessment)
    else:
        # Keep the requirement list, but every assessment of a different text
        # is stale, including missing clauses and structural observations.
        template = next(iter(assessments.values()), {})
        item.nli_after = deepcopy(template)
        item.nli_after.update(
            verdict="unknown", issues=[],
            warnings=["Wybrana wersja nie ma zapisanej oceny NLI; wymaga ponownej weryfikacji klauzul."],
        )
        total = 0
        for section in item.nli_after.get("sections", []):
            section["verdict"] = "unknown"
            for clause in section.get("clauses", []):
                clause.update(verdict="unknown", assessment="not_rechecked")
                total += 1
        item.nli_after["coverage"] = {
            "total": total, "checked": 0, "model_checked": 0,
            "structural_checked": 0, "unknown": total,
        }
    item.warnings.extend(item.nli_after.get("warnings", []))
    item.warnings.extend(f"NLI: {issue}" for issue in item.nli_after.get("issues", []))
    if item.nli_after.get("verdict") != "entailed" or item.nli_after.get("coverage", {}).get("unknown", 0):
        item.needs_review = True
        if item.readiness_status == "ready":
            item.readiness_status = "ready_with_warnings"


def apply_review(plan: dict, accepted: list[str], *, source_file: Path | None = None,
                 output: Path | None = None, report: Path | None = None, pdf: Path | None = None):
    """Apply selections and remeasure locally. Does not contact a model."""
    from humanize_pl.flow import humanize
    from humanize_pl.flows.base import FlowSettings, attach_pdf_report, summarise
    from humanize_pl.flows.docx_flow import _write_docx
    from humanize_pl.general import GeneralOptions
    from humanize_pl.io.docx_structure import document_text, load_document

    text, drafts = select_review(plan, accepted)
    if output and output.suffix.lower() not in {".txt", ".docx", ".xlsx"}:
        raise ValueError("Wynik przeglądu zapisujemy jako .txt, .docx albo .xlsx.")
    sources = [source_file] if source_file else []
    ensure_distinct_paths(sources, [p for p in (output, report, pdf) if p is not None])
    if source_file and source_file.suffix.lower() == ".xlsx":
        actual = _workbook_source(source_file, plan)
        if actual != plan["source"]:
            raise ValueError("Komórka źródłowa zmieniła się od przygotowania propozycji.")
    elif source_file:
        actual = document_text(load_document(source_file)) if source_file.suffix.lower() == ".docx" else source_file.read_text(encoding="utf-8")
        if actual != plan["source"]:
            raise ValueError("Wybrany plik nie jest źródłem tego przeglądu.")
    if output and output.suffix.lower() == ".docx" and (not source_file or source_file.suffix.lower() != ".docx"):
        raise ValueError("Zapis DOCX wymaga oryginalnego pliku DOCX.")
    if output and output.suffix.lower() == ".xlsx" and (not source_file or source_file.suffix.lower() != ".xlsx"):
        raise ValueError("Zapis XLSX wymaga oryginalnego pliku XLSX.")
    analysis = plan.get("analysis_settings") or {}
    settings = FlowSettings(rewrite=False, document_type=DocumentType(plan["document_type"]),
                            general_options=GeneralOptions(**(plan.get("general_options") or {})),
                            check_completeness=analysis.get("check_completeness", True),
                            blueprint=analysis.get("blueprint"),
                            style_profile=Path(analysis["style_profile"]) if analysis.get("style_profile") else None,
                            format_policy=FormatPolicy.preserve)
    baseline = humanize(plan["source"], settings=settings, pdf=False)
    result = humanize(text, settings=settings, pdf=False)
    item, before = result.outcomes[0], baseline.outcomes[0]
    item.signal_interpretable = item.signal_interpretable and before.signal_interpretable
    for name in ("signal_before", "findings_before", "family_counts_before", "metrics_before",
                 "blueprint_before", "tone_before", "style_compliance_before", "artifacts_before", "editorial_before"):
        setattr(item, name, getattr(before, name))
    selected = [row for row in plan["proposals"] if row["id"] in accepted]
    item.applied_changes = [{**row, "issue": "review_accepted", "risk": None} for row in selected if row["kind"] == "edit"]
    item.changes_applied = len(item.applied_changes)
    item.examples = item.applied_changes[:4]
    item.drafted_sections = drafts
    item.requested_operations["editing"] = True
    item.requested_operations["drafting"] = any(row["kind"] == "draft" for row in plan["proposals"])
    _restore_nli(plan, text, item)
    if drafts:
        item.needs_review = True
        item.readiness_status = "ready_with_warnings" if item.readiness_status == "ready" else item.readiness_status
        item.warnings.append("Wybrane dopiski modelu wymagają osobnego zatwierdzenia przez prawnika.")
    item.notes.append("Decyzje przeglądu zastosowano do źródła; pomiary dotyczą wybranej wersji. Nie ponawiano zapytań do modelu ani NLI klauzul.")
    item.review = plan
    if output:
        if output.suffix.lower() == ".docx":
            formatting = _write_docx(source_file, output, text, settings=settings,
                                     document_type=settings.document_type, drafts=drafts_from_payload(drafts))
            if not formatting.inventory_preserved:
                raise ValueError("Kontrola DOCX odrzuciła wybrane zmiany; zapisano kopię źródła.")
            item.formatting = formatting.to_json()
            item.warnings.extend(formatting.warnings + formatting.issues)
        elif output.suffix.lower() == ".xlsx":
            _write_workbook_review(source_file, output, plan, text)
        else:
            write_text_atomic(output, text, sources=sources)
    if item.warnings and item.readiness_status == "ready":
        item.readiness_status = "ready_with_warnings"
    decisions = {row["id"]: "accepted" if row["id"] in accepted else "rejected" for row in plan["proposals"]}
    result.payload.update({"review": plan, "review_decisions": decisions,
                           "summary": summarise([item]), "documents": [{**item.to_json(), "applied_changes": item.applied_changes}]})
    result.payload['settings']['rewrite'] = True
    result.payload['settings']['nli'] = item.requested_operations.get('nli', False)
    result.payload['layers']['review_selection'] = True
    result.changes_applied, result.applied_changes = item.changes_applied, item.applied_changes
    result.signal_before, result.signal_delta = item.signal_before, item.signal_delta
    result.drafted_sections, result.operations = drafts, item.operation_results()
    result.readiness_status, result.needs_review = item.readiness_status, item.needs_review
    result.formatting = item.formatting
    result.nli = item.nli
    result.output_path = output
    if pdf:
        attach_pdf_report(result.payload, pdf)
        result.pdf_report = Path(result.payload['pdf_report']) if result.payload.get('pdf_report') else None
    if report:
        write_text_atomic(report, json.dumps(result.payload, ensure_ascii=False, indent=2) + "\n", sources=sources)
        result.report_path = report
    return result


def _workbook_source(source: Path, plan: dict) -> str:
    import openpyxl

    location = plan.get("workbook") or {}
    if not location or not isinstance(location.get("row"), int) or not isinstance(location.get("column"), int):
        raise ValueError("Raport nie wskazuje komórki źródłowej.")
    if location['row'] < 1 or location['column'] < 1:
        raise ValueError("Nieprawidłowa komórka źródłowa.")
    for values_only in (True, False):
        workbook = openpyxl.load_workbook(source, data_only=values_only)
        try:
            value = workbook[location['sheet']].cell(location['row'], location['column']).value
            if value is not None:
                return str(value).strip()
        finally:
            workbook.close()
    return ""


def _write_workbook_review(source: Path, target: Path, plan: dict, text: str) -> None:
    import openpyxl

    if _workbook_source(source, plan) != plan['source']:
        raise ValueError("Komórka źródłowa nie pasuje do propozycji.")
    workbook = openpyxl.load_workbook(source)
    try:
        location = plan['workbook']
        sheet = workbook[location['sheet']]
        column = sheet.max_column + 1
        original = sheet.cell(location['row'], location['column'])
        result = sheet.cell(location['row'], column, text)
        result._style = copy(original._style)
        if location.get('header_row'):
            sheet.cell(location['header_row'], column, 'Tekst po przeglądzie')
        sheet.column_dimensions[result.column_letter].width = sheet.column_dimensions[original.column_letter].width
        with atomic_output(target, sources=[source]) as staged:
            workbook.save(staged)
    finally:
        workbook.close()


def load_reviews(path: Path) -> list[dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("schema") == 1 and "proposals" in payload:
        return [payload]
    if payload.get("review"):
        return [payload["review"]]
    plans = []
    for row in payload.get("documents", payload.get("rows", [])):
        if not row.get("review", {}).get("available"):
            continue
        plan = deepcopy(row["review"])
        # Older reports kept NLI beside the plan. Preserve it when loading
        # those reports instead of silently treating verification as disabled.
        if "nli" not in plan:
            output, _ = select_review(plan, [p["id"] for p in plan["proposals"]])
            plan["nli"] = _nli_context(
                plan["source"], output,
                requested=row.get("requested_operations", {}).get("nli", payload.get("settings", {}).get("nli", False)),
                before=row.get("nli_before", {}), after=row.get("nli_after", row.get("nli", {})),
            )
        plans.append(plan)
    return plans


def review_markdown(plan: dict) -> str:
    from html import escape

    parts = ["Zmiany zależne w jednym akapicie stanowią jedną propozycję. Zaznaczone propozycje zostaną zachowane; pozostałe przywrócą źródło."]
    for row in plan.get("proposals", []):
        parts.append(f"\n### {row['id']}\n")
        if row.get("previous"):
            parts.append("Przed: " + escape(row["previous"]))
        parts.append("<pre>" + escape(row["before"] or "(dopisek)") + "</pre>\n↓\n<pre>" + escape(row["after"]) + "</pre>")
        if row.get("following"):
            parts.append("Dalej: " + escape(row["following"]))
    return "\n\n".join(parts)
