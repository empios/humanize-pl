"""Traces of the tool rather than the style: markdown, chatbot asides, fields.

Every pattern here was measured against 2396 court judgments before it was
written, and the negative tests below are the human uses that measurement
turned up.
"""

from __future__ import annotations

from humanize_pl.artifacts import find_artifacts, strip_markup

MODEL_OUTPUT = (
    "Poniżej znajdziesz gotowy wzór umowy.\n"
    "# UMOWA O ŚWIADCZENIE USŁUG\n"
    "---\n"
    "**§ 1. Przedmiot umowy**\n"
    "Wykonawca wykona usługi w terminie do [data].\n"
    "| Usługa | Cena |\n"
    "|---|---|\n"
    "Podpis: ....................\n"
    "Jeśli chcesz, mogę przygotować wersję dwujęzyczną."
)


def test_each_trace_of_model_output_is_found():
    counts = find_artifacts(MODEL_OUTPUT).counts()

    assert counts == {
        "chatbot_frame": 2,
        "markdown_heading": 1,
        "markdown_rule": 1,
        "markdown_bold": 1,
        "placeholder": 1,
        "markdown_table": 1,
    }


def test_fields_to_fill_include_dotted_blanks():
    """A dotted blank is not an AI trace - judgments quoting forms have them -
    but a document carrying one is not ready to send."""
    report = find_artifacts(MODEL_OUTPUT)

    assert report.fields == ["[data]", "...................."]
    assert "placeholder" in report.counts()
    assert report.to_json()["fields"] == 2


def test_what_judgments_use_is_not_a_trace():
    """The human uses the measurement found: hyphen lists (756 of 2396),
    citations "[w:]", recording times, words inserted into a quotation, "***"
    between sections, a line of underscores for a signature."""
    human = (
        "Sąd zważył, co następuje:\n"
        "- powód nie wykazał szkody,\n"
        "- pozwany nie zaprzeczył.\n"
        "Zob. T. Nowak [w:] Komentarz, s. 12.\n"
        "Świadek zeznał [adn.: 01:14:08-01:32:57 - k. 143].\n"
        "Organ odmówił [wyrejestrowania] pojazdu.\n"
        "***\n"
        "____________________"
    )

    assert find_artifacts(human).counts() == {}


def test_markup_is_removed_line_for_line():
    """A DOCX paragraph is a line and the inventory guard counts them, so a
    rule line becomes an empty line rather than disappearing."""
    text, changes = strip_markup(MODEL_OUTPUT)
    lines = text.split("\n")

    assert len(lines) == len(MODEL_OUTPUT.split("\n"))
    assert lines[1] == "UMOWA O ŚWIADCZENIE USŁUG"
    assert lines[2] == ""
    assert lines[3] == "§ 1. Przedmiot umowy"
    # Content decisions are left to a person.
    assert lines[0].startswith("Poniżej znajdziesz")
    assert lines[6] == "|---|---|"
    assert "[data]" in lines[4]
    assert {change["issue"] for change in changes} == {"markdown"}


def test_protected_lines_are_not_touched():
    text, changes = strip_markup("**a**\n**b**", protected={1})

    assert text == "a\n**b**"
    assert [change["paragraph_index"] for change in changes] == [0]


def test_an_em_dash_is_still_a_dash_and_a_rule_line_is_not():
    from humanize_pl.detect import detect_document

    def dashes(text: str) -> int:
        diagnosis = detect_document(text, calibrate_against_default=False)
        return sum(row.count for row in diagnosis.families if row.family == "typography_artifact")

    assert dashes("Sąd ustalił—wbrew twierdzeniom—że powództwo jest zasadne.") == 2
    assert dashes("Umowa.\n---\nDalej.\n|---|---|") == 0


def test_the_flow_strips_markup_and_holds_back_what_it_cannot_decide():
    from humanize_pl.config import Engine, Mode
    from humanize_pl.document import ReadinessStatus
    from humanize_pl.flows import FlowSettings, run_all_layers

    outcome, _verdict = run_all_layers(
        MODEL_OUTPUT,
        name="umowa.txt",
        settings=FlowSettings(mode=Mode.standard, engine=Engine.basic, draft_missing=False),
    )

    assert "**" not in outcome.text_out and "# " not in outcome.text_out
    assert len(outcome.text_out.split("\n")) == len(MODEL_OUTPUT.split("\n"))
    after = outcome.artifacts_after["counts"]
    assert not {"markdown_bold", "markdown_heading", "markdown_rule"} & set(after)
    assert outcome.artifacts_before["counts"]["markdown_bold"] == 1
    assert any("czatbota" in warning for warning in outcome.warnings)
    assert any("2 pola do uzupełnienia" in warning for warning in outcome.warnings)
    # The stub also lacks the sections a contract owes, which fails it on its
    # own; what matters here is that the traces alone never leave it ready.
    assert outcome.needs_review is True
    assert outcome.readiness_status != ReadinessStatus.ready.value


def test_a_docx_loses_its_markdown_and_keeps_its_paragraphs(tmp_path):
    from docx import Document

    from humanize_pl.config import Engine, Mode
    from humanize_pl.flows import FlowSettings, run_docx_flow

    source = tmp_path / "in"
    source.mkdir()
    document = Document()
    for line in ("# WEZWANIE DO ZAPŁATY", "---", "**Wzywam** do zapłaty kwoty 1 200 zł."):
        document.add_paragraph(line)
    document.save(str(source / "wezwanie.docx"))

    payload = run_docx_flow(
        source,
        tmp_path / "out",
        settings=FlowSettings(mode=Mode.standard, engine=Engine.basic, draft_missing=False),
        pdf=False,
    )

    row = payload["documents"][0]
    written = Document(str(tmp_path / "out" / "wezwanie_humanized.docx"))
    texts = [paragraph.text for paragraph in written.paragraphs]
    assert row["formatting"]["inventory_preserved"] is True
    assert texts[0] == "WEZWANIE DO ZAPŁATY"
    assert texts[1] == ""
    assert "**" not in texts[2]


def test_blank_lines_between_numbered_points_survive_the_rules():
    """The enumeration protector's leading whitespace swallowed the newline of a blank
    line, and the spacing cleanup then stripped it: a real model contract lost
    53 of its 92 blank lines. A .txt output has no other paragraph breaks."""
    from humanize_pl.config import Engine
    from humanize_pl.core import create_humanizer_session
    from humanize_pl.safety.protectors import protect_text

    text = "Zawarta pomiędzy:\n\n1. Spółką X\n\n§ 1. Przedmiot umowy\n\n1. Zamawiający zleca usługi.\n\n2. Usługi obejmują:"

    assert "\n\n" in protect_text(text).text.split("__PROTECTED")[0]
    assert create_humanizer_session(engine=Engine.basic).humanize(text).text == text


def test_a_field_counts_the_same_with_or_without_its_bold():
    """Measured on the corpus: fields rose 590 -> 619 across a rewrite that
    added none, because "**……… zł**" hid its blank until the bold came off."""
    assert find_artifacts("Kwota: **……… zł**").fields == find_artifacts("Kwota: ……… zł").fields
    assert len(find_artifacts("Kwota: **……… zł**").fields) == 1


def test_the_chatbots_other_asides_are_found_too():
    """Each from a real model document the first pattern missed."""
    asides = (
        "Poniżej gotowy wzór wniosku. Uzupełnij pola w nawiasach kwadratowych.",
        "Poniżej projekt pozwu cywilnego o zapłatę wynagrodzenia z umowy o dzieło.",
        "Opinia ma charakter ogólny i nie stanowi porady prawnej w konkretnej sprawie.",
        "*Dokument ma charakter wzorcowy. Przed podpisaniem należy dostosować treść.*",
        "Chcesz żebym dopasował ten regulamin do Twojej konkretnej branży?",
        "**Uwaga praktyczna:** jeżeli faktury mają różne terminy płatności, wskaż je.",
    )
    for line in asides:
        assert find_artifacts(line).counts().get("chatbot_frame"), line


def test_words_that_only_begin_like_an_aside_are_not_one():
    """Judgment text the first version of the wider pattern caught."""
    for line in (
        "Ławy fundamentowe posadowiono 80 cm poniżej projektowanego poziomu terenu.",
        "Świadkowie wskazali na zeznania wzajemnie się uzupełniające.",
    ):
        assert not find_artifacts(line).counts().get("chatbot_frame"), line


def test_a_second_models_asides_are_found_too():
    """Read off qwen-local, the pattern found Bielik's asides in 11 of 32
    documents; these are the shapes that made it 25."""
    for line in (
        "Oto przykładowy wniosek o udostępnienie informacji publicznej:",
        "Oto kompletna polityka prywatności dla aplikacji mobilnej.",
        "Pamiętaj, że to jest tylko przykładowa umowa i może wymagać dostosowania.",
        "Ta polityka prywatności może być dostosowana do konkretnych potrzeb platformy.",
        "W razie wątpliwości warto skonsultować się z prawnikiem.",
    ):
        assert find_artifacts(line).counts().get("chatbot_frame"), line
    # A party who consulted a lawyer, in a judgment, is not an aside.
    assert not find_artifacts("Powód skonsultował się z prawnikiem.").counts()
