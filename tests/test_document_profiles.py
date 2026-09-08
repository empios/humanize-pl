from __future__ import annotations

import json
from pathlib import Path

from docx import Document
import pytest
from typer.testing import CliRunner

from humanize_pl.config import Engine
from humanize_pl.document import (
    DocumentType,
    RewriteBackend,
    StyleProfile,
    build_style_profile,
    classify_document,
)
from humanize_pl.flows import FlowSettings, run_all_layers
from humanize_pl.flows.cli import app


def _write_docx(path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        (
            "Szanowni Państwo, informujemy o ryzyku i rekomendujemy dalsze kroki.",
            DocumentType.client_communication,
        ),
        (
            "Niniejsza umowa stanowi, że Strona zobowiązuje się zapłacić wynagrodzenie.",
            DocumentType.contract,
        ),
        (
            "Sąd, sygn. akt I C 1/26. Wnoszę o wydanie postanowienia. Uzasadnienie.",
            DocumentType.filing_official,
        ),
    ],
)
def test_document_type_classifier(text: str, expected: DocumentType) -> None:
    guess = classify_document(text)
    assert guess.document_type == expected
    assert 0 <= guess.confidence <= 1


def test_profile_stores_only_short_anonymized_examples_and_style_data(tmp_path) -> None:
    samples = tmp_path / "samples"
    samples.mkdir()
    for index in range(5):
        _write_docx(
            samples / f"sample-{index}.docx",
            f"Jan Kowalski zapłacił {5000 + index} zł. "
            "Dalsza część zatwierdzonego pisma opisuje rekomendację dla klienta.",
        )
    guide = tmp_path / "guide.yaml"
    guide.write_text(
        "preferred_terms:\n  dokonać zapłaty: zapłacić\n"
        "forbidden_phrases:\n  - warto podkreślić\n"
        "voice:\n  - rzeczowy\n",
        encoding="utf-8",
    )
    output = tmp_path / "profile"
    profile = build_style_profile(
        samples,
        output,
        name="Kancelaria testowa",
        document_type=DocumentType.client_communication,
        style_guide=guide,
    )

    loaded = StyleProfile.load(output)
    assert loaded.name == profile.name
    assert loaded.document_count == 5
    assert loaded.preferred_terms["dokonać zapłaty"] == "zapłacić"
    serialized = json.dumps(loaded.to_json(), ensure_ascii=False)
    assert "Jan Kowalski" not in serialized
    assert "5000" not in serialized
    assert all(len(example) <= 260 for example in loaded.anonymized_examples)


def test_profile_command_builds_bundle(tmp_path) -> None:
    samples = tmp_path / "samples"
    samples.mkdir()
    for index in range(5):
        _write_docx(samples / f"{index}.docx", "Informujemy klienta o dalszych krokach.")
    output = tmp_path / "profile"
    result = CliRunner().invoke(
        app,
        [
            "profile",
            str(samples),
            "--name",
            "Test",
            "--document-type",
            "client_communication",
            "-o",
            str(output),
        ],
    )
    assert result.exit_code == 0, result.stdout
    assert (output / "profile.json").exists()


def test_hybrid_backend_falls_back_to_rules_with_warning(monkeypatch, tmp_path) -> None:
    monkeypatch.chdir(tmp_path)
    for key in (
        "HUMANIZE_PL_LLM_BASE_URL",
        "HUMANIZE_PL_LLM_MODEL",
        "HUMANIZE_PL_LLM_API_KEY",
    ):
        monkeypatch.delenv(key, raising=False)
    settings = FlowSettings(
        engine=Engine.basic,
        rewrite_backend=RewriteBackend.hybrid,
    )
    outcome, _verdict = run_all_layers(
        "Warto podkreślić, że Pracownik wykonuje obowiązek.",
        name="test",
        settings=settings,
    )
    assert outcome.llm["status"] == "unavailable"
    assert outcome.readiness_status == "ready_with_warnings"
    assert any("HUMANIZE_PL_LLM_BASE_URL" in warning for warning in outcome.warnings)


def test_require_llm_turns_missing_configuration_into_error(monkeypatch, tmp_path) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("HUMANIZE_PL_LLM_BASE_URL", raising=False)
    monkeypatch.delenv("HUMANIZE_PL_LLM_MODEL", raising=False)
    with pytest.raises(RuntimeError):
        run_all_layers(
            "Warto podkreślić, że Pracownik wykonuje obowiązek.",
            name="test",
            settings=FlowSettings(
                engine=Engine.basic,
                rewrite_backend=RewriteBackend.hybrid,
                require_llm=True,
            ),
        )


def _office(tmp_path, texts, *, name="Kancelaria", document_type=None):
    import docx as pydocx

    from humanize_pl.document import DocumentType, build_style_profile

    source = tmp_path / "in"
    source.mkdir(parents=True, exist_ok=True)
    for index, text in enumerate(texts):
        document = pydocx.Document()
        for paragraph in [row for row in text.split("\n") if row.strip()]:
            document.add_paragraph(paragraph)
        document.save(source / f"d{index:02d}.docx")
    return build_style_profile(
        source_directory=source,
        output_directory=tmp_path / "out",
        name=name,
        document_type=document_type or DocumentType.contract,
    )


HUMAN_CONTRACT = (
    "§ {n}. Przedmiot umowy\n"
    "Wykonawca wykona dokumentację techniczną węzła cieplnego przy ul. Kolejowej 4. "
    "Zakres obejmuje inwentaryzację, projekt i nadzór. Termin: 30 września.\n"
    "§ {m}. Wynagrodzenie\n"
    "Za wykonanie przedmiotu umowy Zamawiający zapłaci 42 000 zł netto, płatne po "
    "odbiorze. Faktura płatna w 21 dni.\n"
)


def test_office_profile_carries_a_full_measurement(tmp_path) -> None:
    """Three ad-hoc numbers are not a baseline; the real measurement is.

    An office profile has to be comparable with the public reference profiles,
    otherwise a firm's own documents cannot serve as the yardstick for its own
    work.
    """
    profile = _office(
        tmp_path, [HUMAN_CONTRACT.format(n=i, m=i + 1) for i in range(1, 13, 2)]
    )
    reference = profile.reference_profile()
    assert reference is not None
    assert reference.name.startswith("office:")
    assert reference.genre == "contract"
    assert reference.sentence_words.p50 > 0
    # Every known family is measured, at zero where the office produces none.
    # A baseline that lists only what it happened to see stops calibrating the
    # rest, so clean writing would yield the blindest profile.
    assert len(reference.family_rates) >= 10
    assert reference.family_rates["discourse_frame"].p95 == 0.0


def test_a_thin_office_profile_is_reported_as_indicative(tmp_path) -> None:
    """Percentiles need volume. Six documents give a picture, not a threshold."""
    profile = _office(tmp_path, [HUMAN_CONTRACT.format(n=i, m=i + 1) for i in range(1, 13, 2)])
    assert profile.document_count < 30
    assert profile.reference_is_indicative is True


def test_documents_that_look_ai_written_are_refused_as_a_baseline(tmp_path) -> None:
    """The failure that disables the tool without anyone noticing.

    A baseline built from AI-assisted drafts teaches the detector that machine
    style is normal. It then reports low scores on exactly the documents it
    was bought to catch - and a low score reads as good news.
    """
    ai_texts = [
        (Path("docs_tests/ai_generated") / name).read_text(encoding="utf-8")
        for name in (
            "ai_legal_01_umowa_uslug.txt",
            "ai_legal_04_regulamin_platformy.txt",
            "ai_legal_08_polityka_rodo.txt",
            "ai_legal_09_umowa_it.txt",
            "ai_legal_11_regulamin_sklepu.txt",
            "ai_legal_13_polityka_aplikacji.txt",
        )
    ]
    profile = _office(tmp_path, ai_texts)
    assert profile.warnings, "profil z dokumentów AI musi ostrzegać"
    assert "ślady pisania przez AI" in profile.warnings[0]

    clean = _office(
        tmp_path / "clean",
        [HUMAN_CONTRACT.format(n=i, m=i + 1) for i in range(1, 13, 2)],
    )
    assert clean.warnings == []


def test_the_office_warning_reaches_every_run_that_uses_the_profile(tmp_path) -> None:
    from humanize_pl.config import Engine, Mode
    from humanize_pl.document import DocumentType
    from humanize_pl.flows.base import FlowSettings, run_all_layers

    ai_texts = [
        (Path("docs_tests/ai_generated") / name).read_text(encoding="utf-8")
        for name in (
            "ai_legal_01_umowa_uslug.txt",
            "ai_legal_04_regulamin_platformy.txt",
            "ai_legal_08_polityka_rodo.txt",
            "ai_legal_09_umowa_it.txt",
            "ai_legal_11_regulamin_sklepu.txt",
            "ai_legal_13_polityka_aplikacji.txt",
        )
    ]
    _office(tmp_path, ai_texts, name="Kancelaria AI")
    outcome, _verdict = run_all_layers(
        HUMAN_CONTRACT.format(n=1, m=2),
        name="umowa.docx",
        settings=FlowSettings(
            mode=Mode.standard,
            engine=Engine.basic,
            rewrite=False,
            document_type=DocumentType.contract,
            style_profile=tmp_path / "out" / "profile.json",
        ),
    )
    assert any("ślady pisania przez AI" in warning for warning in outcome.warnings)
    assert outcome.calibration_status.startswith("calibrated:office:")
