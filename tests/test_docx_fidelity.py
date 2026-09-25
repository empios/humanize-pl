"""Realistic synthetic DOCX structures, with deliberate same-count corruption."""

import json
from copy import deepcopy
from dataclasses import replace
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from humanize_pl.document import FormatPolicy
from humanize_pl.flow import humanize
from humanize_pl.flows import FlowSettings, run_all_layers, run_docx_flow
from humanize_pl.io.docx_structure import (
    document_text,
    inventory_docx,
    iter_text_units,
    new_paragraph_like,
    replace_unit_text,
    save_with_inventory_guard,
)


def _element(tag, text=None, **attributes):
    node = OxmlElement(f"w:{tag}")
    if text is not None:
        node.text = text
    for key, value in attributes.items():
        node.set(qn(f"w:{key}"), value)
    return node


def _note(document, kind="footnote"):
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    xml = (f'<w:{kind}s xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f'<w:{kind} w:id="1"><w:p><w:r><w:t>Przypis: zakres dokumentacji obejmuje załącznik.</w:t>'
           f'</w:r></w:p></w:{kind}></w:{kind}s>')
    part = Part(PackURI(f'/word/{kind}s.xml'),
                f'application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml',
                xml.encode('utf-8'), document.part.package)
    document.part.relate_to(part, f'http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s')
    document.add_paragraph("Zakres dokumentacji").add_run()._r.append(_element(f"{kind}Reference", id="1"))


def _contract(path):
    document = Document()
    document.add_heading("Umowa o przygotowanie dokumentacji", 0)
    document.add_paragraph("§ 1. Przedmiot umowy")
    paragraph = document.add_paragraph()
    paragraph.add_run("Warto podkreślić, że ")
    paragraph.add_run("Wykonawca przygotuje dokumentację.").italic = True
    document.add_paragraph("Zakres obejmuje dokumentację wykonawczą.", style="List Number")
    document.add_paragraph("Odbiór następuje protokołem.", style="List Number")
    paragraph = document.add_paragraph("Szczegóły: ")
    paragraph._p.append(_element("bookmarkStart", id="10", name="zakres"))
    paragraph.add_run("załącznik nr 1")
    paragraph._p.append(_element("bookmarkEnd", id="10"))
    paragraph = document.add_paragraph("Zobacz ")
    link = _element("hyperlink")
    link.set(qn("r:id"), document.part.relate_to('https://example.test/zalacznik',
             'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
    run = _element("r")
    run.append(_element("t", "załącznik"))
    link.append(run)
    paragraph._p.append(link)
    paragraph = document.add_paragraph("Odesłanie do zakresu: ")
    field = _element("fldSimple", instr="REF zakres \\h")
    run = _element("r")
    run.append(_element("t", "załącznik nr 1"))
    field.append(run)
    paragraph._p.append(field)
    _note(document)
    _note(document, "endnote")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Harmonogram i odbiór"
    cell = table.cell(1, 0)
    cell.text = "Przed tabelą zagnieżdżoną."
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Etap 1 — dokumentacja."
    cell.add_paragraph("Po tabeli zagnieżdżonej.")
    table.cell(1, 1).text = "Należy zauważyć, że termin wynosi 14 dni."
    first = document.add_paragraph("Początek komentarza.").runs[0]
    document.add_paragraph("Warto podkreślić, że środkowy akapit wymaga uzgodnienia.")
    last = document.add_paragraph("Koniec komentarza.").runs[0]
    document.add_comment([first, last], text="Propozycja do uzgodnienia.", author="Recenzent", initials="R")
    paragraph = document.add_paragraph()
    insertion = _element("ins", id="1", author="Recenzent")
    run = _element("r")
    run.append(_element("t", "Należy zauważyć, że treść jest w śledzonej zmianie."))
    insertion.append(run)
    paragraph._p.append(insertion)
    deletion = _element("del", id="2", author="Recenzent")
    deleted_run = _element("r")
    deleted_run.append(_element("delText", "Usunięta propozycja do porównania."))
    deletion.append(deleted_run)
    paragraph._p.append(deletion)
    control = _element("sdt")
    content = _element("sdtContent")
    wrapped = document.add_paragraph("Warto podkreślić, że to treść kontrolki.")
    content.append(wrapped._p)
    control.append(content)
    document.element.body.insert(len(document.element.body) - 1, control)
    paragraph = document.add_paragraph("Podpis wykonawcy")
    paragraph.add_run("\t…\nPodpis zamawiającego\t…")
    document.sections[0].header.paragraphs[0].text = "Projekt umowy / dokumentacja"
    document.sections[0].footer.paragraphs[0].text = "Egzemplarz do uzgodnienia"
    document.save(path)
    return document


def test_contract_preserves_protected_content_and_reports_its_scope(tmp_path):
    source, target = tmp_path / "source.docx", tmp_path / "out.docx"
    original = _contract(source)
    result = humanize(source, output=target, check_completeness=False, pdf=False)
    assert result.ok, result.payload["documents"][0].get("error")
    written = Document(target)
    assert result.changed
    assert result.text == document_text(written)
    assert "Wykonawca przygotuje dokumentację." in result.text
    assert "Warto podkreślić, że Wykonawca" not in result.text
    assert "Warto podkreślić, że środkowy akapit" in result.text
    assert "Warto podkreślić, że to treść kontrolki." in result.text
    assert "Podpis wykonawcy\t…\u2028Podpis zamawiającego\t…" in result.text
    units = {unit.location: unit for unit in iter_text_units(written)}
    for unit in iter_text_units(original):
        if unit.protected:
            assert units[unit.location].paragraph._p.xml == unit.paragraph._p.xml
    assert inventory_docx(source).structural_differences(inventory_docx(target)) == []
    assert result.formatting["skipped_units"]
    assert "word/footnotes.xml" in result.formatting["excluded_parts"]
    with ZipFile(source) as before, ZipFile(target) as after:
        for name in ("word/footnotes.xml", "word/endnotes.xml", "word/comments.xml", "word/numbering.xml"):
            assert before.read(name) == after.read(name)
    text = result.text
    assert text.index("Przed tabelą") < text.index("Etap 1") < text.index("Po tabeli")
    assert text.count("Harmonogram i odbiór") == 1


def test_disjoint_edits_preserve_the_run_of_unchanged_emphasis():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Wstęp: ")
    paragraph.add_run("ważne zdanie").bold = True
    paragraph.add_run(". Zbędny koniec.")
    unit = next(iter_text_units(document))
    replace_unit_text(unit, "Ważne: ważne zdanie.")
    assert paragraph.text == "Ważne: ważne zdanie."
    assert paragraph.runs[1].text == "ważne zdanie"
    assert paragraph.runs[1].bold is True


def test_capitalisation_after_intro_removal_uses_the_body_run_format():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Warto podkreślić, że ").bold = True
    paragraph.add_run("termin wynosi 14 dni.").italic = True
    replace_unit_text(next(iter_text_units(document)), "Termin wynosi 14 dni.")
    assert paragraph.runs[0].text == ""
    assert paragraph.runs[1].text == "Termin wynosi 14 dni."
    assert paragraph.runs[1].italic


@pytest.mark.parametrize("replacement", ["XabCD", "aXbCD", "abXCD", "abCDX", "aD", ""])
def test_insertions_and_deletions_have_exact_text_without_flattening_runs(replacement):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("ab").italic = True
    paragraph.add_run("CD").bold = True
    replace_unit_text(next(iter_text_units(document)), replacement)
    assert paragraph.text == replacement
    assert paragraph.runs[0].italic and paragraph.runs[1].bold


@pytest.mark.parametrize("feature", ["hyperlink", "link_target", "bookmark", "field", "footnote",
                                      "comment", "revision", "numbering", "style_numbering", "merge", "format"])
def test_same_count_corruption_is_reverted_to_exact_source(tmp_path, feature):
    source, target = tmp_path / "source.docx", tmp_path / "out.docx"
    document = _contract(source)
    body = document.element.body
    if feature == "hyperlink":
        body.xpath(".//w:hyperlink//w:t")[0].text = "Inny tekst odsyłacza"
    elif feature == "link_target":
        rel_id = body.xpath(".//w:hyperlink")[0].get(qn("r:id"))
        document.part.rels[rel_id]._target = "https://example.test/inny-cel"
    elif feature == "bookmark":
        body.xpath(".//w:bookmarkStart")[0].set(qn("w:name"), "inny_cel")
    elif feature == "field":
        body.xpath(".//w:fldSimple")[0].set(qn("w:instr"), "REF inny_cel")
    elif feature == "footnote":
        body.xpath(".//w:footnoteReference")[0].set(qn("w:id"), "2")
    elif feature == "comment":
        next(iter(document.comments)).paragraphs[0].runs[0].text = "Zmieniony komentarz."
    elif feature == "revision":
        body.xpath(".//w:ins")[0].set(qn("w:author"), "Inny autor")
    elif feature == "numbering":
        document.part.numbering_part.element.xpath(".//w:start")[0].set(qn("w:val"), "9")
    elif feature == "style_numbering":
        document.styles["List Number"].element.xpath(".//w:numId")[0].set(qn("w:val"), "9")
    elif feature == "merge":
        body.xpath(".//w:gridSpan")[0].set(qn("w:val"), "3")
    else:
        document.paragraphs[2].runs[1].italic = False
    differences = save_with_inventory_guard(document, source, target)
    assert differences
    assert target.read_bytes() == source.read_bytes()


def test_drafted_paragraph_disables_numbering_inherited_from_its_style():
    document = Document()
    paragraph = document.add_paragraph("Pozycja", style="List Number")
    created = new_paragraph_like(paragraph, "Dopisana sekcja")
    assert created.xpath("./w:pPr/w:numPr/w:numId")[0].get(qn("w:val")) == "0"


def test_multi_paragraph_field_protects_the_middle_paragraph():
    document = Document()
    document.add_paragraph().add_run()._r.append(_element("fldChar", fldCharType="begin"))
    middle = document.add_paragraph("Warto podkreślić, że to wynik pola.")
    document.add_paragraph().add_run()._r.append(_element("fldChar", fldCharType="end"))
    unit = next(unit for unit in iter_text_units(document) if unit.paragraph._p is middle._p)
    assert unit.protected
    assert "wieloakapitowe pole dokumentu" in unit.protection_reasons


def test_reverted_output_has_source_metrics_gate_and_no_stale_success(tmp_path, monkeypatch):
    from humanize_pl.flows import docx_flow

    source, target = tmp_path / "source.docx", tmp_path / "out.docx"
    document = Document()
    document.add_paragraph("Warto podkreślić, że termin wynosi 14 dni.")
    document.save(source)
    settings = FlowSettings(check_completeness=False)
    expected, expected_gate = run_all_layers(document_text(document), name=source.name,
                                             settings=replace(settings, rewrite=False))
    actual_save = docx_flow.save_with_inventory_guard
    calls = []

    def corrupt(doc, *args, **kwargs):
        calls.append(True)
        doc.paragraphs[0].runs[0].bold = True
        return actual_save(doc, *args, **kwargs)

    monkeypatch.setattr(docx_flow, "save_with_inventory_guard", corrupt)
    for resume in (False, True):
        payload = run_docx_flow(source, target, settings=settings, resume=resume, pdf=False)
        result = payload["documents"][0]
        assert target.read_bytes() == source.read_bytes()
        assert result["text_out"] == document_text(Document(target))
        assert result["signal_after"] == result["signal_before"] == expected.signal_after
        assert result["changes_applied"] == 0
        assert result["metrics_after"] == expected.metrics_after
        assert result["family_counts_after"] == expected.family_counts_after
        assert result["unresolved_findings"] == expected.unresolved_findings
        assert result["compliance"] == expected.compliance
        assert result["constraints"] == expected.constraints
        assert result["legal_sensitive_check"] == expected.legal_sensitive_check
        assert result["operations"]["editing"]["status"] == "not_saved"
        assert result["formatting"]["fixes"] == []
        assert result["needs_review"]
        assert payload["summary"]["mean_signal_delta"] == 0
        detail = json.loads((tmp_path / "details" / "source.json").read_text(encoding="utf-8"))
        assert detail["gate"] == expected_gate.to_json()
        assert detail["resume"]["reusable"] is False
    assert len(calls) == 2


@pytest.mark.parametrize("source_nli", [{}, {"verdict": "unknown", "coverage": {"unknown": 2}}])
def test_rollback_never_reuses_nli_of_the_discarded_text(source_nli):
    from humanize_pl.flows.docx_flow import _measure_restored_source

    settings = FlowSettings(nli=True)
    before, _ = run_all_layers("Warto podkreślić, że termin wynosi 14 dni.", name="umowa",
                                settings=replace(settings, nli=False))
    before.nli_before = deepcopy(source_nli)
    before.requested_operations = settings.requested_operations()
    before.nli_after = {"verdict": "entailed", "coverage": {"unknown": 0}}
    restored, _ = _measure_restored_source("Warto podkreślić, że termin wynosi 14 dni.", before, settings, None)
    assert restored.nli_after == source_nli
    assert restored.operation_results()["completeness"]["status"] == "incomplete"


def test_read_only_docx_does_not_claim_a_stale_target(tmp_path):
    source, target = tmp_path / "source.docx", tmp_path / "out.docx"
    document = Document()
    document.add_paragraph("Treść do analizy.")
    document.save(source)
    target.write_bytes(b"earlier unrelated result")
    result = humanize(source, output=target, no_rewrite=True, pdf=False)
    assert result.output_path is None
    assert result.text == "Treść do analizy."
    assert target.read_bytes() == b"earlier unrelated result"


def test_audit_without_renderer_is_explicit(tmp_path, monkeypatch):
    from humanize_pl.docx_quality import audit_document, render_and_audit

    path = tmp_path / "source.docx"
    document = _contract(path)
    report = audit_document(document, path, policy=FormatPolicy.audit)
    monkeypatch.setattr("humanize_pl.docx_quality._tool", lambda _name: None)
    render_and_audit(path, report)
    assert not report.rendered
    assert report.text_complete is None
    assert any("pominięto kontrolę renderowania" in warning for warning in report.warnings)


@pytest.mark.parametrize("anchor_kind", ["table", "comment"])
def test_drafts_do_not_enter_a_table_cell_or_a_protected_range(tmp_path, anchor_kind):
    from humanize_pl.document import DocumentType
    from humanize_pl.drafting import DraftedSection, insert_drafts
    from humanize_pl.flows.docx_flow import _write_docx

    source, target = tmp_path / "source.docx", tmp_path / "out.docx"
    document = Document()
    if anchor_kind == "table":
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Treść komórki."
    else:
        run = document.add_paragraph("Treść z komentarzem.").runs[0]
        document.add_comment([run], text="Uwaga", author="R")
    document.save(source)
    draft = DraftedSection("test", "sekcja", "Propozycja.", (), 0)
    text = insert_drafts(document_text(document), [draft])
    report = _write_docx(source, target, text, settings=FlowSettings(),
                         document_type=DocumentType.contract, drafts=[draft])
    assert not report.inventory_preserved
    assert "miejsce dopisania" in report.inventory_differences[0]
    assert target.read_bytes() == source.read_bytes()


def test_drafting_next_to_a_numbered_paragraph_preserves_existing_numbering(tmp_path):
    from humanize_pl.document import DocumentType
    from humanize_pl.drafting import DraftedSection, insert_drafts
    from humanize_pl.flows.docx_flow import _write_docx

    source, target = tmp_path / "source.docx", tmp_path / "out.docx"
    document = Document()
    document.add_paragraph("Pierwsza pozycja.", style="List Number")
    document.add_paragraph("Druga pozycja.", style="List Number")
    document.save(source)
    draft = DraftedSection("test", "sekcja", "Propozycja.", (), 0)
    text = insert_drafts(document_text(document), [draft])
    report = _write_docx(source, target, text, settings=FlowSettings(),
                         document_type=DocumentType.contract, drafts=[draft])
    assert report.inventory_preserved, report.inventory_differences
    written = Document(target)
    assert [p.text for p in written.paragraphs] == ["Pierwsza pozycja.", "Propozycja.", "Druga pozycja."]
    assert written.paragraphs[1]._p.xpath("./w:pPr/w:numPr/w:numId")[0].get(qn("w:val")) == "0"
