"""Missing sections supplied inside the flow, for text and for DOCX.

`test_drafting.py` covers what a draft may contain and where it goes. These
cover what the flow does with it: the document that leaves carries the
clause, the structure is measured on that document, nothing reports it as
ready, and a DOCX gains exactly the paragraphs the flow announced.
"""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from humanize_pl.config import Engine, Mode
from humanize_pl.document import ReadinessStatus
from humanize_pl.flows import FlowSettings, run_all_layers, run_docx_flow
from humanize_pl.llm import LlmSettings

pytest.importorskip("docx")

FIXTURE = Path("docs_tests/ai_generated/ai_legal_01_umowa_uslug.txt")
SETTINGS = FlowSettings(mode=Mode.standard, engine=Engine.basic, draft_missing=True)

# One clause per section the fixture lacks, written the way the prompt asks:
# a blank where a fact belongs, and the words the skeleton recognises.
ANSWERS = {
    "rozwiązanie i wypowiedzenie umowy": (
        "Każda ze Stron może dokonać wypowiedzenia umowy z zachowaniem … "
        "okresu wypowiedzenia, w formie pisemnej pod rygorem nieważności."
    ),
    "postanowienia końcowe": (
        "W sprawach nieuregulowanych niniejszą umową stosuje się przepisy "
        "Kodeksu cywilnego. Zmiany umowy wymagają formy pisemnej."
    ),
    "podpisy stron": "Zamawiający: … Wykonawca: …",
}


class FakeModel:
    """Answers each drafting request by the section it names."""

    def __init__(self, answers: dict[str, str] = ANSWERS, presence: dict[str, dict] | None = None) -> None:
        self.settings = LlmSettings("https://model.test/v1", "fake")
        self.answers = answers
        # What the model says when asked whether a section is already there;
        # by default every section the fixture lacks really is missing.
        self.presence = presence or {}
        self.metadata = SimpleNamespace(to_report=lambda: {"backend": "fake"})

    def complete_json(self, messages, *, max_tokens=None) -> dict:
        prompt = messages[-1]["content"]
        for label, answer in self.presence.items():
            if f"Część: {label}." in prompt:
                return answer
        return {"obecna": False, "cytat": ""}

    def complete_text(self, messages, *, max_tokens=None, temperature=None) -> str:
        prompt = messages[-1]["content"]
        for label, answer in self.answers.items():
            if f"Brakująca sekcja: {label}." in prompt:
                return answer
        raise AssertionError(f"nieoczekiwane zapytanie: {prompt[:120]}")

    def close(self) -> None:
        pass


def test_the_flow_supplies_missing_sections_and_measures_the_result():
    text = FIXTURE.read_text(encoding="utf-8")

    outcome, _verdict = run_all_layers(
        text, name="umowa.txt", settings=SETTINGS, rewriter=FakeModel()
    )

    assert outcome.blueprint_before["blocking"] is True
    assert [row["label_pl"] for row in outcome.drafted_sections] == list(ANSWERS)
    for answer in ANSWERS.values():
        assert answer in outcome.text_out
    # Measured on the document that leaves, so the gaps are closed...
    assert outcome.blueprint_after["missing_required"] == []
    assert outcome.blueprint_after["numbering_issues"] == []
    # ...and still nothing a machine wrote reports as ready to send.
    assert outcome.readiness_status == ReadinessStatus.ready_with_warnings.value
    assert outcome.needs_review is True
    assert any("Dopisano sekcję" in warning for warning in outcome.warnings)


def test_drafted_clauses_do_not_trip_the_rewrite_inventory_check():
    """The check guards the rewrite of existing text. A drafted clause
    repeating "Zamawiający" is new text with its own gate, not a rewrite that
    changed a party name."""
    text = FIXTURE.read_text(encoding="utf-8")

    outcome, _verdict = run_all_layers(
        text, name="umowa.txt", settings=SETTINGS, rewriter=FakeModel()
    )

    assert outcome.legal_sensitive_check["passed"] is True


def test_without_a_model_the_gap_is_reported_not_silently_left(monkeypatch):
    monkeypatch.setattr("humanize_pl.flows.base.prepare_llm", lambda settings: (None, []))
    text = FIXTURE.read_text(encoding="utf-8")

    outcome, _verdict = run_all_layers(text, name="umowa.txt", settings=SETTINGS)

    assert outcome.drafted_sections == []
    assert any("nie dopisano" in warning for warning in outcome.warnings)
    assert outcome.readiness_status == ReadinessStatus.failed.value


def test_drafting_can_be_switched_off():
    from dataclasses import replace

    text = FIXTURE.read_text(encoding="utf-8")

    outcome, _verdict = run_all_layers(
        text,
        name="umowa.txt",
        settings=replace(SETTINGS, draft_missing=False),
        rewriter=FakeModel(),
    )

    assert outcome.drafted_sections == []
    assert outcome.readiness_status == ReadinessStatus.failed.value


def _write_fixture_docx(path: Path) -> list[str]:
    from docx import Document

    lines = [
        line for line in FIXTURE.read_text(encoding="utf-8").split("\n") if line.strip()
    ]
    document = Document()
    for line in lines:
        paragraph = document.add_paragraph(line)
        if line.startswith("§"):
            paragraph.runs[0].bold = True
    document.save(str(path))
    return lines


def test_a_docx_gains_exactly_the_announced_paragraphs(tmp_path, monkeypatch):
    from docx import Document

    from humanize_pl.flows import docx_flow

    source = tmp_path / "in"
    source.mkdir()
    original = _write_fixture_docx(source / "umowa.docx")
    monkeypatch.setattr(docx_flow, "prepare_llm", lambda settings: (FakeModel(), []))

    payload = run_docx_flow(source, tmp_path / "out", settings=SETTINGS, pdf=False)

    row = payload["documents"][0]
    written = Document(str(tmp_path / "out" / "umowa_humanized.docx"))
    paragraphs = [p.text for p in written.paragraphs if p.text.strip()]
    drafted_lines = sum(2 if r["heading"] else 1 for r in row["drafted_sections"])

    assert row["formatting"]["inventory_preserved"] is True
    assert all(r["inserted"] for r in row["drafted_sections"])
    assert len(paragraphs) == len(original) + drafted_lines
    for answer in ANSWERS.values():
        assert answer in paragraphs
    # After the last unit nothing needs renumbering, so the clauses take the
    # next numbers; the signature block stays outside the numbering.
    assert [r["heading"] for r in row["drafted_sections"]] == [
        "§ 6. Rozwiązanie i wypowiedzenie umowy",
        "§ 7. Postanowienia końcowe",
        "",
    ]
    # A new heading takes the look of the unit heading it follows.
    headings = [p for p in written.paragraphs if p.text in {"§ 6. Rozwiązanie i wypowiedzenie umowy", "§ 7. Postanowienia końcowe"}]
    assert len(headings) == 2 and all(p.runs[0].bold for p in headings)
    body = next(p for p in written.paragraphs if p.text == ANSWERS["postanowienia końcowe"])
    assert not body.runs[0].bold
    assert row["readiness_status"] == ReadinessStatus.ready_with_warnings.value


def test_a_withdrawn_docx_save_reports_the_drafts_as_proposals(tmp_path, monkeypatch):
    """When the guard restores the source, the file has no drafted clause, and
    the report must neither claim one nor report the structure as fixed."""
    from humanize_pl.flows import docx_flow

    source = tmp_path / "in"
    source.mkdir()
    _write_fixture_docx(source / "umowa.docx")
    monkeypatch.setattr(docx_flow, "prepare_llm", lambda settings: (FakeModel(), []))
    # Announce one paragraph fewer than is created.
    real = docx_flow._insert_drafted_paragraphs
    monkeypatch.setattr(
        docx_flow,
        "_insert_drafted_paragraphs",
        lambda units, drafts: real(units, drafts) - 1,
    )

    payload = run_docx_flow(source, tmp_path / "out", settings=SETTINGS, pdf=False)

    row = payload["documents"][0]
    assert row["formatting"]["inventory_preserved"] is False
    assert row["drafted_sections"] and not any(r["inserted"] for r in row["drafted_sections"])
    assert row["blueprint_after"]["blocking"] is True
    assert row["readiness_status"] == ReadinessStatus.failed.value
    detail = json.loads(
        (tmp_path / "out" / "details" / "umowa.json").read_text(encoding="utf-8")
    )
    assert detail["drafted_sections"][0]["inserted"] is False
    assert row["operations"]["drafting"]["status"] == "not_saved"
    assert row["operations"]["drafting"]["sections_added"] == 0
    assert row["operations"]["drafting"]["requires_review"] is True


def test_the_docx_command_can_leave_missing_sections_unwritten(monkeypatch, tmp_path):
    """Only `run` had the switch: from `docx`, drafting could be turned off
    only by turning the hosted model off altogether."""
    from typer.testing import CliRunner

    from humanize_pl import cli

    seen = {}

    def fake_flow(folder, output, *, settings, **_kwargs):
        seen["draft_missing"] = settings.draft_missing
        raise RuntimeError("stop here")

    monkeypatch.setattr(cli, "run_docx_flow", fake_flow)
    CliRunner().invoke(cli.app, ["docx", str(tmp_path), "--no-draft-missing", "--no-pdf"])
    assert seen["draft_missing"] is False
    CliRunner().invoke(cli.app, ["docx", str(tmp_path), "--no-pdf"])
    assert seen["draft_missing"] is False
    CliRunner().invoke(cli.app, ["docx", str(tmp_path), "--draft-missing", "--no-pdf"])
    assert seen["draft_missing"] is True


def test_a_section_the_model_finds_in_other_words_is_not_drafted():
    """The firm writes "Każda ze stron może wypowiedzieć…" where the pattern
    wants "rozwiązanie umowy": asked first, the model points at it."""
    text = FIXTURE.read_text(encoding="utf-8")
    quote = next(line.strip() for line in text.split("\n") if len(line.strip()) > 40)
    model = FakeModel(presence={"rozwiązanie i wypowiedzenie umowy": {"obecna": True, "cytat": quote}})

    outcome, _ = run_all_layers(text, name="umowa.txt", settings=SETTINGS, rewriter=model)

    drafted = [row["label_pl"] for row in outcome.drafted_sections]
    assert "rozwiązanie i wypowiedzenie umowy" not in drafted
    assert outcome.sections_found_by_model["rozwiązanie i wypowiedzenie umowy"] == quote
    assert any("model wskazuje, że już jest" in warning for warning in outcome.warnings)


def test_an_invented_quote_neither_drafts_nor_clears_the_section():
    text = FIXTURE.read_text(encoding="utf-8")
    model = FakeModel(presence={"rozwiązanie i wypowiedzenie umowy": {
        "obecna": True, "cytat": "Każda ze stron może wypowiedzieć umowę z miesięcznym wyprzedzeniem."}})

    outcome, _ = run_all_layers(text, name="umowa.txt", settings=SETTINGS, rewriter=model)

    assert "rozwiązanie i wypowiedzenie umowy" not in [row["label_pl"] for row in outcome.drafted_sections]
    assert "rozwiązanie i wypowiedzenie umowy" not in outcome.sections_found_by_model
    assert any("cytat, którego nie ma w dokumencie" in warning for warning in outcome.warnings)
    # Still missing as far as anyone can tell, so the document still fails.
    assert outcome.readiness_status == ReadinessStatus.failed.value


def test_an_answer_without_a_verdict_drafts_nothing():
    text = FIXTURE.read_text(encoding="utf-8")
    model = FakeModel(presence={label: {"cytat": ""} for label in ANSWERS})

    outcome, _ = run_all_layers(text, name="umowa.txt", settings=SETTINGS, rewriter=model)

    assert outcome.drafted_sections == []
    assert outcome.text_out.count("Kodeksu cywilnego") == text.count("Kodeksu cywilnego")


def test_drafting_is_opt_in_even_with_an_available_model():
    class NoCalls(FakeModel):
        def complete_json(self, *_args, **_kwargs):
            pytest.fail("Default settings must not ask a model about missing sections")

        def complete_text(self, *_args, **_kwargs):
            pytest.fail("Default settings must not draft")

    outcome, _ = run_all_layers(
        FIXTURE.read_text(encoding="utf-8"), name="umowa", settings=FlowSettings(), rewriter=NoCalls(),
    )
    assert outcome.drafted_sections == []
    assert outcome.operation_results()["drafting"]["status"] == "disabled"


def test_draft_only_preserves_existing_prose_and_has_separate_results(monkeypatch):
    from humanize_pl.flow import humanize

    monkeypatch.setattr("humanize_pl.flow.prepare_llm", lambda settings: (FakeModel(), []))
    source = FIXTURE.read_text(encoding="utf-8")
    result = humanize(source, no_rewrite=True, draft_missing=True)

    assert result.changed
    assert result.changes_applied == 0
    assert result.applied_changes == []
    assert result.operations["editing"]["status"] == "disabled"
    assert result.operations["completeness"]["status"] == "completed"
    assert result.operations["drafting"] == {
        "requested": True, "status": "added", "sections_added": 3, "requires_review": True,
    }
    assert [r["text"] for r in result.drafted_sections] == list(ANSWERS.values())
    # Existing paragraphs remain verbatim and in their original order.
    old_lines = source.splitlines()
    assert [line for line in result.text.splitlines() if line in old_lines] == old_lines


def test_draft_only_docx_is_saved_and_resume_checks_its_fingerprint(tmp_path, monkeypatch):
    from dataclasses import replace

    from docx import Document

    from humanize_pl.flows import docx_flow

    source = tmp_path / "umowa.docx"
    original = _write_fixture_docx(source)
    before = source.read_bytes()
    target = tmp_path / "out" / "umowa_humanized.docx"
    monkeypatch.setattr(docx_flow, "prepare_llm", lambda settings: (FakeModel(), []))
    settings = replace(SETTINGS, rewrite=False)
    payload = run_docx_flow(source, target, settings=settings, pdf=False)
    paragraphs = [p.text for p in Document(target).paragraphs]
    assert source.read_bytes() == before
    assert [p for p in paragraphs if p in original] == original
    assert all(answer in paragraphs for answer in ANSWERS.values())
    assert payload["documents"][0]["changes_applied"] == 0
    assert payload["documents"][0]["operations"]["drafting"]["sections_added"] == 3

    # A modified output must be regenerated, even with editing disabled.
    target.write_bytes(b"interrupted output")
    resumed = run_docx_flow(source, target, settings=settings, pdf=False, resume=True)
    assert resumed["summary"]["failed"] == 0
    assert all(answer in [p.text for p in Document(target).paragraphs] for answer in ANSWERS.values())


def test_text_ui_shows_full_drafts_and_scope(monkeypatch):
    from humanize_pl.ui.app import run_text

    monkeypatch.setattr("humanize_pl.flow.prepare_llm", lambda settings: (FakeModel(), []))
    _text, summary, changes, _gate = run_text(
        FIXTURE.read_text(encoding="utf-8"), None,
        "standard", "basic", "auto", "rules", "preserve", False,
        False, False, False, False, False, None, None, "(brak)", False, True, "legal", True,
    )
    assert "Redakcja językowa: wyłączona" in summary
    assert "Kontrola kompletności: wykonana" in summary
    assert "Sekcje dodane: 3" in summary
    assert "wymaga przeglądu prawnika" in changes
    assert all(answer in changes for answer in ANSWERS.values())
