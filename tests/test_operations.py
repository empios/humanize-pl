"""Independent controls must affect execution, saved outputs and reports."""

from pathlib import Path
from types import SimpleNamespace

import pytest
from typer.testing import CliRunner

from humanize_pl import cli
from humanize_pl.flow import humanize
from humanize_pl.flows import FlowSettings
from humanize_pl.flows.base import ItemOutcome, prepare_llm
from humanize_pl.reports.operations import operation_lines

FIXTURE = Path("docs_tests/ai_generated/ai_legal_01_umowa_uslug.txt")


def test_completeness_only_never_loads_a_model_or_changes_text(monkeypatch):
    def forbidden(*_args, **_kwargs):
        pytest.fail("Completeness patterns need neither a model nor an editing session")

    monkeypatch.setattr("humanize_pl.flows.base.LlmSettings.from_environment", forbidden)
    monkeypatch.setattr(FlowSettings, "session", forbidden)
    source = FIXTURE.read_text(encoding="utf-8")
    result = humanize(source, no_rewrite=True, rewrite_backend="hybrid")
    assert result.text == source
    assert not result.changed
    assert result.blueprint["missing_required"]
    assert result.operations["completeness"]["status"] == "completed"
    assert result.operations["editing"]["status"] == "disabled"
    assert result.operations["drafting"]["status"] == "disabled"


def test_editing_without_completeness_does_not_resolve_or_check_blueprints(monkeypatch):
    def forbidden(*_args):
        pytest.fail("Disabled completeness must not load or check a blueprint")

    monkeypatch.setattr("humanize_pl.flows.base.resolve_blueprint", forbidden)
    monkeypatch.setattr("humanize_pl.flows.base.blueprint_for", forbidden)
    monkeypatch.setattr("humanize_pl.flows.base.check", forbidden)
    result = humanize("Warto podkreślić, że termin wynosi 14 dni.", check_completeness=False)
    assert result.text == "Termin wynosi 14 dni."
    assert not result.blueprint["checked"]
    assert result.operations["editing"]["changes"] == 1
    assert result.operations["completeness"]["status"] == "disabled"


@pytest.mark.parametrize("option", ["nli", "draft_missing"])
def test_dependent_operations_require_completeness(option):
    with pytest.raises(ValueError, match="włączonej kontroli kompletności"):
        humanize("Tekst.", check_completeness=False, **{option: True})


def test_drafting_uses_the_configured_model_without_enabling_rewriting(monkeypatch):
    client = SimpleNamespace(probe=lambda: True)
    config = object()
    monkeypatch.setattr("humanize_pl.flows.base.LlmSettings.from_environment", lambda _path: config)
    monkeypatch.setattr("humanize_pl.flows.base.OpenAICompatibleRewriter", lambda settings: client if settings is config else None)
    model, warnings = prepare_llm(FlowSettings(rewrite=False, draft_missing=True))
    assert model is client
    assert warnings == []


@pytest.mark.parametrize("command", ["run", "docx", "xlsx"])
def test_cli_forwards_completeness_and_drafting_defaults(command, tmp_path, monkeypatch):
    seen = {}

    def stop(*_args, **kwargs):
        settings = kwargs.get("settings")
        seen.update(vars(settings) if settings else kwargs)
        raise ValueError("stopped before processing")

    if command == "run":
        monkeypatch.setattr(cli, "humanize", stop)
        args = ["run", "Tekst."]
    elif command == "docx":
        monkeypatch.setattr(cli, "run_docx_flow", stop)
        args = ["docx", str(tmp_path)]
    else:
        monkeypatch.setattr(cli, "run_xlsx_flow", stop)
        source = tmp_path / "in.xlsx"
        source.touch()
        args = ["xlsx", str(source), "--column", "A"]
    result = CliRunner().invoke(cli.app, [*args, "--no-check-completeness", "--no-pdf"])
    assert seen, result.output
    assert seen["check_completeness"] is False
    assert seen["draft_missing"] is False


def test_xlsx_rejects_explicit_drafting_before_writing(tmp_path):
    source = tmp_path / "in.xlsx"
    source.write_bytes(b"source not opened")
    target = tmp_path / "out.xlsx"
    with pytest.raises(ValueError, match="Dopisywanie sekcji nie jest obsługiwane"):
        humanize(source, output=target, column="A", draft_missing=True)
    assert not target.exists()
    assert source.read_bytes() == b"source not opened"


def test_xlsx_contains_operation_results_without_edits(tmp_path):
    from openpyxl import Workbook, load_workbook

    source = tmp_path / "in.xlsx"
    workbook = Workbook()
    workbook.active.append(["Treść"])
    workbook.active.append(["Zwykły tekst."])
    workbook.save(source)
    workbook.close()
    target = tmp_path / "out.xlsx"
    humanize(source, output=target, column="Treść", no_rewrite=True, check_completeness=False)
    written = load_workbook(target)
    sheet = written.active
    column = next(cell.column for cell in sheet[1] if cell.value == "Zakres i wyniki czynności")
    assert "Kontrola kompletności: wyłączona" in sheet.cell(2, column).value
    assert "Redakcja językowa: wyłączona" in sheet.cell(2, column).value
    written.close()


def test_report_distinguishes_unavailable_and_incomplete_checks():
    item = ItemOutcome(name="umowa", requested_operations=FlowSettings(nli=True).requested_operations())
    assert item.operation_results()["completeness"]["status"] == "incomplete"
    item.requested_operations["nli"] = False
    assert item.operation_results()["completeness"]["status"] == "unavailable"
    item.status = "failed"
    assert "Kontrola kompletności: błąd" in "\n".join(operation_lines([item.to_json()]))
    assert operation_lines([{"name": "old report"}]) == []


def test_drafting_failure_does_not_replace_per_document_edit_counters(monkeypatch):
    from humanize_pl.document import RewriteBackend
    from humanize_pl.flows import run_all_layers
    from humanize_pl.llm import LlmBatchMetadata

    metadata = LlmBatchMetadata(model="fake", status="ready", proposals=8, accepted=5, rejected=3)

    def rewrite(text, *_args, **_kwargs):
        metadata.proposals += 1
        metadata.rejected += 1
        return text, [], 1

    def draft(text, *_args, **_kwargs):
        metadata.status = "ready_with_errors"
        return text

    monkeypatch.setattr("humanize_pl.flows.base._rewrite_remaining_with_llm", rewrite)
    monkeypatch.setattr("humanize_pl.flows.base._supply_missing_sections", draft)
    settings = FlowSettings(rewrite_backend=RewriteBackend.hybrid, draft_missing=True)
    outcome, _ = run_all_layers(FIXTURE.read_text(encoding="utf-8"), name="umowa", settings=settings,
                                rewriter=SimpleNamespace(metadata=metadata))
    assert outcome.llm["proposals"] == 1
    assert outcome.llm["accepted"] == 0
    assert outcome.llm["rejected"] == 1
    assert outcome.llm["status"] == "ready_with_errors"
