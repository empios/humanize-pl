"""Tests for the plain-language Polish PDF report."""

from __future__ import annotations

import pytest

from humanize_pl.config import Engine, Mode
from humanize_pl.flows import FlowSettings, run_docx_flow, run_xlsx_flow
from humanize_pl.gate import FAMILY_CONSTRAINTS
from humanize_pl.reports import pdf_pl

reportlab = pytest.importorskip("reportlab")

BASIC = FlowSettings(mode=Mode.standard, engine=Engine.basic)

AI_TEXT = (
    "Podporządkowanie pracownika jest jedną z najważniejszych cech stosunku pracy.\n"
    "Warto wskazać, że podporządkowanie nie oznacza całkowitej zależności. "
    "Kluczowe znaczenie ma tutaj kierownictwo pracodawcy.\n"
    "Z jednej strony umożliwia to organizowanie pracy. Podsumowując, ma to duże znaczenie."
)


def write_docx(path, text: str) -> None:
    from docx import Document

    document = Document()
    for paragraph in text.split("\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph)
    document.save(str(path))


def test_docx_flow_writes_the_pdf_next_to_the_json_report(tmp_path) -> None:
    source = tmp_path / "in"
    source.mkdir()
    write_docx(source / "opinia.docx", AI_TEXT)
    output = tmp_path / "out"

    payload = run_docx_flow(source, output, settings=BASIC)

    report = output / "raport.pdf"
    assert report.exists()
    assert report.read_bytes().startswith(b"%PDF")
    assert payload["pdf_report"] == str(report)
    assert payload["pdf_error"] is None


def test_pdf_can_be_turned_off(tmp_path) -> None:
    source = tmp_path / "in"
    source.mkdir()
    write_docx(source / "opinia.docx", AI_TEXT)
    output = tmp_path / "out"

    payload = run_docx_flow(source, output, settings=BASIC, pdf=False)

    assert not (output / "raport.pdf").exists()
    assert "pdf_report" not in payload


def write_xlsx(path, openpyxl) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.append(["ID", "Odpowiedź AI"])
    workbook.active.append(["1", AI_TEXT])
    workbook.save(str(path))


def test_xlsx_flow_writes_the_pdf_beside_the_workbook(tmp_path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    source = tmp_path / "in.xlsx"
    write_xlsx(source, openpyxl)

    payload = run_xlsx_flow(source, tmp_path / "out.xlsx", column="B", settings=BASIC)

    report = tmp_path / "out_raport.pdf"
    assert report.exists()
    assert payload["pdf_report"] == str(report)


def test_a_missing_dependency_is_reported_not_raised(tmp_path, monkeypatch) -> None:
    """A run whose real work succeeded must not fail over an optional report."""
    from humanize_pl.flows.base import attach_pdf_report

    monkeypatch.setattr(pdf_pl, "pdf_available", lambda: False)
    payload = attach_pdf_report({"flow": "docx", "summary": {}}, tmp_path / "raport.pdf")

    assert payload["pdf_report"] is None
    assert "reportlab" in payload["pdf_error"]
    assert not (tmp_path / "raport.pdf").exists()


def test_a_font_without_polish_glyphs_is_refused(monkeypatch) -> None:
    """Vera ships with reportlab and silently drops ą, ę, ś — better to fail."""
    from pathlib import Path

    import reportlab

    vera = str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf")
    monkeypatch.delenv(pdf_pl.FONT_ENV, raising=False)
    monkeypatch.setattr(pdf_pl, "FONT_CANDIDATES", ((vera, vera),))

    assert not pdf_pl._covers_polish(vera)
    with pytest.raises(pdf_pl.PdfDependencyError, match="polskimi znakami"):
        pdf_pl._resolve_font_paths()


def test_report_renders_without_any_processed_item(tmp_path) -> None:
    target = pdf_pl.write_flow_pdf(
        {"flow": "docx", "settings": {}, "layers": {}, "summary": {}, "documents": []},
        tmp_path / "pusty.pdf",
    )

    assert target.read_bytes().startswith(b"%PDF")


def test_failed_items_reach_the_report(tmp_path) -> None:
    payload = {
        "flow": "docx",
        "settings": {"mode": "standard", "engine": "basic", "rewrite": True},
        "layers": {},
        "summary": {"items": 1, "ok": 0, "failed": 1, "needs_review": 0},
        "documents": [{"name": "zepsuty.docx", "status": "failed", "error": "BadZipFile: x"}],
    }

    target = pdf_pl.write_flow_pdf(payload, tmp_path / "raport.pdf")

    assert target.exists()


def test_every_gate_family_has_a_polish_description() -> None:
    """A new signal family must not reach the client report as a bare slug."""
    missing = sorted(set(FAMILY_CONSTRAINTS) - set(pdf_pl.FAMILY_GLOSSARY))

    assert not missing, f"brak opisu dla rodzin: {missing}"


def test_every_scored_metric_has_a_polish_description() -> None:
    from humanize_pl.detect.engine import _metrics

    measured = set(_metrics([(0, 0, "Pierwsze zdanie testowe ma kilka słów.")], 6,
                            sentences_per_paragraph=[1]))
    described = set(pdf_pl.METRIC_GLOSSARY) | {"words"}

    assert not measured - described


@pytest.mark.parametrize(
    ("count", "expected"),
    [(1, "zmiana"), (2, "zmiany"), (4, "zmiany"), (5, "zmian"), (12, "zmian"), (22, "zmiany")],
)
def test_polish_number_agreement(count: int, expected: str) -> None:
    assert pdf_pl._plural(count, "zmiana", "zmiany", "zmian") == expected


def test_the_report_never_names_a_file(tmp_path) -> None:
    """Odbiorca dostaje rodzaj materiału i numer pozycji, nie nazwę pliku."""
    source = tmp_path / "in"
    source.mkdir()
    write_docx(source / "tajna-opinia-klienta.docx", AI_TEXT)

    payload = run_docx_flow(source, tmp_path / "out", settings=BASIC, pdf=False)
    report = pdf_pl._Report(payload, 400, pdf_pl._styles())

    assert report.material == "DOCX"
    assert pdf_pl._item_label(1, "dokument", payload["documents"][0]) == "Dokument 1"

    rendered = pdf_pl.write_flow_pdf(payload, tmp_path / "raport.pdf").read_bytes()
    assert b"tajna-opinia-klienta" not in rendered


def test_a_spreadsheet_row_keeps_its_number() -> None:
    """Numer wiersza to nie nazwa pliku, więc zostaje: po nim się trafia do danych."""
    assert pdf_pl._item_label(1, "element", {"name": "wiersz 7"}) == "Element 7"


def _detail_payload(items: list[dict]) -> dict:
    return {
        "flow": "xlsx",
        "settings": {"mode": "standard", "engine": "basic", "rewrite": True},
        "layers": {},
        "summary": {"items": len(items), "ok": len(items), "failed": 0, "needs_review": 1},
        "rows": items,
    }


def test_each_element_gets_its_own_changes_and_advice() -> None:
    """Osoba redagująca pyta o jedną pozycję, nie o średnią z paczki."""
    payload = _detail_payload(
        [
            {
                "name": "wiersz 17",
                "status": "ok",
                "signal_before": 0.6,
                "signal_after": 0.4,
                "needs_review": True,
                "changes_applied": 1,
                "examples": [{"before": "Warto wskazać, że tak.", "after": "Tak.",
                              "issue": "legal_ai_style_rewrite"}],
                "constraints": ["Nie kończ akapitem podsumowującym."],
            },
            {
                "name": "wiersz 18",
                "status": "ok",
                "signal_before": 0.1,
                "signal_after": 0.1,
                "needs_review": False,
                "changes_applied": 0,
                "examples": [],
                "constraints": [],
            },
        ]
    )
    report = pdf_pl._Report(payload, 400, pdf_pl._styles())

    detailed, quiet, cut = report._detail_selection()

    assert [position for position, _ in detailed] == [1]
    assert (quiet, cut) == (1, 0)


def test_a_large_batch_keeps_the_pooled_list_as_a_fallback() -> None:
    """Przy 40 pozycjach rozpisujemy część, ale komplet uwag musi gdzieś być."""
    items = [
        {
            "name": f"wiersz {index}",
            "status": "ok",
            "signal_before": 0.6,
            "signal_after": 0.5,
            "needs_review": True,
            "changes_applied": 1,
            "examples": [],
            "constraints": ["Zróżnicuj długość zdań."],
        }
        for index in range(2, 42)
    ]
    report = pdf_pl._Report(_detail_payload(items), 400, pdf_pl._styles())

    detailed, _quiet, cut = report._detail_selection()

    assert len(detailed) == pdf_pl.DETAIL_LIMIT
    assert cut == len(items) - pdf_pl.DETAIL_LIMIT
    assert report._pooled_recommendations(len(detailed), cut)


def test_the_pooled_list_stays_out_when_everything_was_written_out() -> None:
    report = pdf_pl._Report(_detail_payload([]), 400, pdf_pl._styles())

    assert report._pooled_recommendations(3, 0) == []


def test_the_report_avoids_the_phrasing_it_flags() -> None:
    """Raport o maszynowym stylu nie może sam być pisany maszynowo.

    Długie myślniki jako wtrącenie są najbardziej rzucającym się w oczy
    sygnałem; cytaty i przykłady zwrotów są z tej reguły wyłączone.
    """
    import ast
    from pathlib import Path

    quoted = set(pdf_pl.FAMILY_GLOSSARY) | {"example"}
    source = Path(pdf_pl.__file__).read_text(encoding="utf-8")
    polish = [
        node.value
        for node in ast.walk(ast.parse(source))
        if isinstance(node, ast.Constant)
        and isinstance(node.value, str)
        and any(char in node.value for char in "ąćęłńóśźż")
    ]

    offenders = [text for text in polish if "—" in text and text not in quoted]

    assert not offenders, offenders


# --- Regenerating a report from a finished run -------------------------------


LEGACY_FIELDS = (
    "family_counts_before",
    "family_counts_after",
    "metrics_before",
    "metrics_after",
    "findings_after",
    "examples",
)


def make_legacy_run(tmp_path):
    """A completed run whose payload predates the newer measurements."""
    import json

    source = tmp_path / "in"
    source.mkdir()
    write_docx(source / "opinia.docx", AI_TEXT)
    output = tmp_path / "out"
    run_docx_flow(source, output, settings=BASIC, pdf=False)

    report = output / "flow-report.json"
    payload = json.loads(report.read_text(encoding="utf-8"))
    payload["summary"].pop("findings_after", None)
    for item in payload["documents"]:
        for field in LEGACY_FIELDS:
            item.pop(field, None)
    report.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return source, output


def test_report_can_be_rebuilt_from_a_finished_run(tmp_path) -> None:
    """The point of the command: an old batch must not need reprocessing."""
    from humanize_pl.flows import backfill_payload, load_payload, needs_backfill

    _, output = make_legacy_run(tmp_path)

    payload = load_payload(output)
    assert needs_backfill(payload)
    assert backfill_payload(payload) == 1
    assert not needs_backfill(payload)

    item = payload["documents"][0]
    assert item["family_counts_before"]
    assert item["metrics_before"]["sentence_length_cv"] > 0
    # Treść poprawek odtworzona z porównania oryginału z plikiem po redakcji.
    assert item["examples"]
    assert all(pair["before"] != pair["after"] for pair in item["examples"])

    target = pdf_pl.write_flow_pdf(payload, output / "raport.pdf")
    assert target.read_bytes().startswith(b"%PDF")


def test_backfill_keeps_what_the_original_run_reported(tmp_path) -> None:
    from humanize_pl.flows import backfill_payload, load_payload

    _, output = make_legacy_run(tmp_path)
    payload = load_payload(output)
    original = dict(payload["documents"][0])

    backfill_payload(payload)

    for key, value in original.items():
        assert payload["documents"][0][key] == value


def test_backfill_is_skipped_when_the_documents_are_gone(tmp_path) -> None:
    """Missing measurements stay missing rather than being reported as zero."""
    from humanize_pl.flows import backfill_payload, load_payload

    _, output = make_legacy_run(tmp_path)
    payload = load_payload(output)
    payload["input_directory"] = str(tmp_path / "nie-ma")

    assert backfill_payload(payload) == 0
    assert "family_counts_before" not in payload["documents"][0]

    target = pdf_pl.write_flow_pdf(payload, output / "raport.pdf")
    assert target.exists()


def test_loading_points_at_the_missing_report(tmp_path) -> None:
    from humanize_pl.flows import load_payload

    with pytest.raises(FileNotFoundError, match="flow-report.json"):
        load_payload(tmp_path)


def test_xlsx_flow_writes_its_json_report_without_being_asked(tmp_path) -> None:
    """The .docx flow always did; leaving this one silent left nothing to rebuild."""
    import json

    openpyxl = pytest.importorskip("openpyxl")
    source = tmp_path / "in.xlsx"
    write_xlsx(source, openpyxl)

    payload = run_xlsx_flow(source, tmp_path / "out.xlsx", column="B", settings=BASIC)

    report = tmp_path / "out_raport.json"
    assert report.exists()
    assert payload["report_path"] == str(report)

    saved = json.loads(report.read_text(encoding="utf-8"))
    assert saved["rows"][0]["family_counts_before"]


def test_report_can_be_rebuilt_from_a_finished_workbook(tmp_path) -> None:
    """An .xlsx processed before the JSON report existed still has everything."""
    from humanize_pl.flows import payload_from_workbook

    openpyxl = pytest.importorskip("openpyxl")
    source = tmp_path / "in.xlsx"
    write_xlsx(source, openpyxl)
    output = tmp_path / "out.xlsx"
    run_xlsx_flow(source, output, column="B", settings=BASIC, report=False, pdf=False)

    payload = payload_from_workbook(output, column="Odpowiedź AI")

    assert payload["rebuilt"] is True
    row = payload["rows"][0]
    assert row["family_counts_before"]
    assert row["signal_before"] >= row["signal_after"]

    target = pdf_pl.write_flow_pdf(payload, tmp_path / "raport.pdf")
    assert target.read_bytes().startswith(b"%PDF")


def test_a_rebuilt_report_does_not_claim_zero_changes(tmp_path) -> None:
    """"Nie wiadomo" and "nic nie zmieniono" are different answers."""
    from humanize_pl.flows import payload_from_workbook

    openpyxl = pytest.importorskip("openpyxl")
    source = tmp_path / "in.xlsx"
    write_xlsx(source, openpyxl)
    output = tmp_path / "out.xlsx"
    run_xlsx_flow(source, output, column="B", settings=BASIC, report=False, pdf=False)

    payload = payload_from_workbook(output, column="Odpowiedź AI")
    report = pdf_pl._Report(payload, 400, pdf_pl._styles())

    assert not report.changes_known
    assert "nie znalazł" not in report.headline_sentence()


def test_rebuilding_names_the_column_it_could_not_find(tmp_path) -> None:
    from humanize_pl.flows import payload_from_workbook

    openpyxl = pytest.importorskip("openpyxl")
    source = tmp_path / "in.xlsx"
    write_xlsx(source, openpyxl)

    with pytest.raises(ValueError, match="Nie znaleziono kolumny"):
        payload_from_workbook(source, column="Nie ma takiej")


def test_a_multi_step_rewrite_is_shown_as_one_change() -> None:
    """A→B→C is one edit to the reader, not two."""
    collapsed = pdf_pl._collapse_chains(
        [
            {"before": "A", "after": "B", "issue": "x"},
            {"before": "B", "after": "C", "issue": "y"},
            {"before": "D", "after": "E", "issue": "z"},
        ]
    )

    assert [(row["before"], row["after"]) for row in collapsed] == [("A", "C"), ("D", "E")]


def test_the_flow_keeps_examples_of_what_it_changed(tmp_path) -> None:
    source = tmp_path / "in"
    source.mkdir()
    write_docx(source / "opinia.docx", AI_TEXT)

    payload = run_docx_flow(source, tmp_path / "out", settings=BASIC, pdf=False)

    examples = payload["documents"][0]["examples"]
    assert examples
    assert all(row["before"] != row["after"] for row in examples)


def test_xlsx_json_report_can_be_turned_off_or_moved(tmp_path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    source = tmp_path / "in.xlsx"
    write_xlsx(source, openpyxl)

    run_xlsx_flow(source, tmp_path / "a.xlsx", column="B", settings=BASIC, report=False)
    assert not (tmp_path / "a_raport.json").exists()

    run_xlsx_flow(
        source,
        tmp_path / "b.xlsx",
        column="B",
        settings=BASIC,
        report_path=tmp_path / "gdzie-indziej.json",
    )
    assert (tmp_path / "gdzie-indziej.json").exists()
    assert not (tmp_path / "b_raport.json").exists()


def test_examples_are_recovered_by_comparing_the_two_texts() -> None:
    """Stary przebieg nie zapisywał treści poprawek, ale oba teksty ją niosą."""
    from humanize_pl.flows.replay import reconstruct_examples

    before = "Warto wskazać, że termin minął. Kwota wynosi 500 zł."
    after = "Termin minął. Kwota wynosi 500 zł."

    pairs = reconstruct_examples(before, after)

    assert pairs == [
        {"before": "Warto wskazać, że termin minął.", "after": "Termin minął.", "issue": ""}
    ]


def test_an_untouched_text_yields_no_examples() -> None:
    from humanize_pl.flows.replay import reconstruct_examples

    assert reconstruct_examples("Termin minął.", "Termin minął.") == []
