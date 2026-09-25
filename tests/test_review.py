from copy import deepcopy

import pytest
from docx import Document

from humanize_pl import humanize
from humanize_pl.flows.base import ItemOutcome
from humanize_pl.io.docx_structure import document_text, load_document
from humanize_pl.review import apply_review, create_review, select_review


def plan():
    return create_review('Pierwszy tekst.\n\nDrugi tekst.', ItemOutcome(
        name='sample', text_out='Pierwszy akapit.\n\nDrugi akapit.', document_type='general'))


def test_selection_is_reversible_and_uses_original_offsets():
    review = plan()
    ids = [row['id'] for row in review['proposals']]
    assert select_review(review, [ids[1]])[0] == 'Pierwszy tekst.\n\nDrugi akapit.'
    assert select_review(review, ids)[0] == 'Pierwszy akapit.\n\nDrugi akapit.'
    assert select_review(review, [])[0] == review['source']


@pytest.mark.parametrize('mutation', ['hash', 'before', 'index', 'duplicate', 'newline'])
def test_mismatched_or_overlapping_proposals_fail(mutation):
    review = deepcopy(plan())
    row = review['proposals'][0]
    if mutation == 'hash':
        review['source'] += 'zmiana'
    elif mutation == 'before':
        row['before'] = 'Inny tekst.'
    elif mutation == 'index':
        row['paragraph_index'] = -1
    elif mutation == 'duplicate':
        review['proposals'].append(row)
    else:
        row['after'] += '\nNowy akapit.'
    with pytest.raises(ValueError):
        select_review(review, [])


def test_undo_recalculates_actual_text_and_report_without_model(monkeypatch, tmp_path):
    original = 'Warto podkreślić, że ogród zimą potrzebuje mniej pracy, niż się wydaje.'
    result = humanize(original, track='general', general_options={'max_shortening': 50}, pdf=False)
    review = result.outcomes[0].review
    assert review['proposals']
    def fail(*args, **kwargs):
        pytest.fail('Review must not call a model')
    monkeypatch.setattr('humanize_pl.llm.OpenAICompatibleRewriter._completion', fail)
    restored = apply_review(review, [], output=tmp_path / 'review.txt', report=tmp_path / 'review.json')
    measured = humanize(original, track='general', no_rewrite=True, pdf=False)
    assert restored.text == original
    assert restored.changes_applied == 0
    assert restored.signal_after == measured.signal_after
    assert restored.outcomes[0].metrics_after == measured.outcomes[0].metrics_after
    assert set(restored.payload['review_decisions'].values()) == {'rejected'}
    assert (tmp_path / 'review.txt').read_text(encoding='utf-8') == original


def test_docx_review_retains_runs_and_can_undo(tmp_path):
    source, output = tmp_path / 'source.docx', tmp_path / 'result.docx'
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Pierwszy ').bold = True
    paragraph.add_run('tekst.').italic = True
    document.add_paragraph('Drugi tekst.')
    document.save(source)
    before_bytes = source.read_bytes()
    review = create_review(document_text(load_document(source)), ItemOutcome(
        name='test', document_type='general', text_out='Pierwszy akapit.\nDrugi akapit.'))
    first = review['proposals'][0]['id']
    result = apply_review(review, [first], source_file=source, output=output)
    assert result.text == 'Pierwszy akapit.\nDrugi tekst.'
    assert document_text(load_document(output)) == result.text
    assert Document(output).paragraphs[0].runs[0].bold
    assert Document(output).paragraphs[0].runs[1].italic
    apply_review(review, [], source_file=source, output=output)
    assert document_text(load_document(output)) == review['source']
    assert source.read_bytes() == before_bytes


def test_review_refuses_wrong_source_or_source_overwrite(tmp_path):
    source = tmp_path / 'source.txt'
    source.write_text('Inny tekst.', encoding='utf-8')
    with pytest.raises(ValueError):
        apply_review(plan(), [], source_file=source, output=tmp_path / 'result.txt')
    source.write_text(plan()['source'], encoding='utf-8')
    with pytest.raises(ValueError):
        apply_review(plan(), [], source_file=source, output=source)


def test_draft_can_be_rejected_independently_of_edit():
    draft = {'section_id': 'extra', 'label_pl': 'Dodatek', 'text': 'Dopisana treść.', 'after_line': 0, 'rank': 0, 'heading': ''}
    review = create_review('Tekst źródłowy.', ItemOutcome(name='test', document_type='contract',
        text_out='Tekst poprawiony.\nDopisana treść.', drafted_sections=[draft]))
    edit, addition = review['proposals']
    assert select_review(review, [edit['id']])[0] == 'Tekst poprawiony.'
    assert select_review(review, [addition['id']])[0] == 'Tekst źródłowy.\nDopisana treść.'


def test_workbook_review_preserves_source_formulas_and_formatting(tmp_path):
    import openpyxl
    from openpyxl.styles import Font

    source, target = tmp_path / 'source.xlsx', tmp_path / 'selected.xlsx'
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(['Tekst', 'Obliczenie'])
    sheet.append(['Pierwszy tekst.', '=1+2'])
    sheet['A2'].font = Font(bold=True)
    workbook.save(source)
    workbook.close()
    review = create_review('Pierwszy tekst.', ItemOutcome(name='row2', document_type='general', text_out='Pierwszy akapit.'))
    review['workbook'] = {'sheet': 'Sheet', 'row': 2, 'column': 1, 'header_row': 1}
    apply_review(review, [review['proposals'][0]['id']], source_file=source, output=target)
    result = openpyxl.load_workbook(target)
    assert result.active['A2'].value == 'Pierwszy tekst.'
    assert result.active['B2'].value == '=1+2'
    assert result.active['C2'].value == 'Pierwszy akapit.'
    assert result.active['C2'].font.bold
    result.close()


def test_review_keeps_completeness_scope_and_pdf_path(tmp_path):
    from humanize_pl import humanize

    original = humanize('Warto wskazać, że umowa dotyczy obsługi ogrodu.',
                        check_completeness=False, pdf=False)
    review = original.outcomes[0].review
    result = apply_review(review, [], output=tmp_path / 'selected.txt', pdf=tmp_path / 'selected.pdf')
    assert not result.operations['completeness']['requested']
    assert result.pdf_report == tmp_path / 'selected.pdf'
    assert result.pdf_report.is_file()
    assert result.text == review['source']


@pytest.mark.parametrize('verdict', ['missing', 'unknown'])
def test_review_preserves_nli_findings_on_unchanged_text(tmp_path, monkeypatch, verdict):
    import json

    from humanize_pl.review import load_reviews

    blueprint = tmp_path / 'structure.yaml'
    blueprint.write_text('''category: test
label_pl: Test
numbering: paragraph
sections:
  - id: wynagrodzenie
    label_pl: Wynagrodzenie
    matches: [wynagrodzenie]
    expects: [Umowa określa termin zapłaty.]
''', encoding='utf-8')

    class Judge:
        def judge_section(self, *, expected_clauses, **kwargs):
            return [verdict] * len(expected_clauses)

    monkeypatch.setattr('humanize_pl.nli.LlmClauseJudge.from_environment', lambda *a: Judge())
    initial = humanize('UMOWA O ŚWIADCZENIE USŁUG\n§ 1. Wynagrodzenie\nWynagrodzenie wynosi sto złotych netto.',
                       document_type='contract', blueprint=blueprint, nli=True, no_rewrite=True, pdf=False)
    def fail(*args, **kwargs):
        pytest.fail('Review must reuse NLI without calling a model')
    monkeypatch.setattr('humanize_pl.nli.LlmClauseJudge.from_environment', fail)

    # Both new plans and old full reports must retain the assessment.
    for legacy in (False, True):
        payload = deepcopy(initial.payload)
        if legacy:
            payload['documents'][0]['review'].pop('nli')
        report = tmp_path / 'report.json'
        report.write_text(json.dumps(payload), encoding='utf-8')
        selected = apply_review(load_reviews(report)[0], [])
        assert selected.text == initial.text
        assert selected.nli == initial.nli
        assert selected.needs_review
        assert selected.readiness_status != 'ready'
        assert selected.operations['completeness']['nli_requested']
        assert selected.payload['settings']['nli']
        assert selected.warnings == initial.warnings


def test_partial_selection_requires_new_nli_but_known_versions_reuse_their_assessment(monkeypatch):
    from humanize_pl.flows.base import FlowSettings
    from humanize_pl.nli import ClauseCheck, NliReport, SectionCheck

    def report(verdict):
        return NliReport(category='test', sections=[SectionCheck(
            'section', 'Treść', verdict, (ClauseCheck('Wymaganie.', verdict),),
        )]).to_json()

    original = 'Pierwszy tekst.\nDrugi tekst.'
    outcome = ItemOutcome(name='test', document_type='contract',
        text_out='Pierwszy akapit.\nDrugi akapit.',
        requested_operations={'nli': True}, nli_before=report('missing'), nli_after=report('entailed'))
    review = create_review(original, outcome, settings=FlowSettings(nli=True))
    snapshot = deepcopy(review)
    ids = [row['id'] for row in review['proposals']]
    def fail(*args, **kwargs):
        pytest.fail('Selection must not call a model')
    monkeypatch.setattr('humanize_pl.nli.LlmClauseJudge.from_environment', fail)

    for accepted, expected in [([], 'missing'), (ids, 'entailed'), ([ids[0]], 'unknown')]:
        result = apply_review(review, accepted)
        assert result.nli['verdict'] == expected
        assert result.outcomes[0].nli_before == outcome.nli_before
        if expected == 'unknown':
            assert result.nli['coverage']['unknown'] == 1
            assert result.nli['coverage']['checked'] == 0
            assert result.operations['completeness']['status'] == 'incomplete'
            assert result.needs_review and result.readiness_status != 'ready'
    assert review == snapshot


def test_review_without_saved_nli_stays_unverified():
    review = create_review('Treść do sprawdzenia.', ItemOutcome(
        name='test', document_type='contract', text_out='Treść do sprawdzenia.',
        requested_operations={'nli': True}))
    result = apply_review(review, [])
    assert result.nli['verdict'] == 'unknown'
    assert result.needs_review
    assert result.operations['completeness']['status'] == 'incomplete'
